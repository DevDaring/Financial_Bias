"""
Practitioner pilot (Improvement_Plan WS5). Reads only the Cases sheet answer columns; the Participant tab (role,
experience) is never opened. Scores each answer against the evidence and compares the three conditions:
T template, G narrative if B1F (GPT-OSS) released it else template, R narrative if ReCheck-P (GPT-OSS) released it
else template.

    python analyze_practitioner.py -> Submission/locked/practitioner_analysis.json
"""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import verify as V  # noqa: E402
from build_pack_m1 import eo_from  # noqa: E402
from evidence import EvidenceType  # noqa: E402
from score_m1 import SUB, yn  # noqa: E402

D = SUB / "packs" / "Annotated" / "IDBI_Annotator_Checks" / "Returned"


def minutes(a, b):
    try:
        def m(x):
            x = str(x).strip()
            hh, mm = re.match(r"(\d{1,2})[:.](\d{2})", x).groups()
            return int(hh) * 60 + int(mm)
        d = m(b) - m(a)
        return d if 0 < d <= 30 else None          # implausible spans are treated as missing
    except Exception:  # noqa: BLE001
        return None


def main() -> int:
    from openpyxl import load_workbook
    key = {json.loads(l)["case_id"]: json.loads(l) for l in open(SUB / "practitioner_key_v1.jsonl")}
    recs = {json.loads(l)["item_id"]: json.loads(l) for l in open(SUB / "locked" / "fresh_natural_v1.jsonl")}
    by = defaultdict(lambda: defaultdict(list))
    n = 0
    for f in sorted(D.rglob("PRACTITIONER_task_P*.xlsx")):
        wb = load_workbook(f, read_only=True, data_only=True)
        rows = list(wb["Cases"].iter_rows(values_only=True))
        head = [str(c) for c in rows[0]]
        for r in rows[1:]:
            d = dict(zip(head, r))
            k = key.get(str(d.get("case_id")))
            if not k:
                continue
            n += 1
            eo = eo_from(recs[k["item_id"]]["evidence"])
            pdv = 100 * eo.prediction.calibrated_pd
            at = max(eo.by_type(EvidenceType.MODEL_ATTRIBUTION), key=lambda a: abs(a.contribution or 0))
            c = k["condition"]
            try:
                q1 = float(re.findall(r"-?\d+(?:\.\d+)?", str(d.get("Q1_PD_VALUE")))[0])
                by[c]["pd_ok"].append(abs(q1 - pdv) <= 0.15 or abs(q1 - round(pdv)) < 0.51 and abs(q1 - pdv) <= 0.55)
            except Exception:  # noqa: BLE001
                by[c]["pd_ok"].append(False)
            names = [V._canon(x.lower()) for x in V._names(at, eo.dataset)] + [str(at.feature_group).lower()]
            ans = V._canon(str(d.get("Q2_STRONGEST_DRIVER") or "").lower())
            by[c]["driver_ok"].append(any(nm and (nm in ans or ans in nm) for nm in names if len(ans) >= 3))
            dirn = str(d.get("Q2_DIRECTION") or "").strip().lower()
            by[c]["direction_ok"].append(("rais" in dirn) == (at.direction == "increases_risk") and dirn != "")
            q3 = yn(d.get("Q3_UNSUPPORTED (Y/N)"))
            if k["shown_wrong"]:
                by[c]["detected_when_wrong"].append(q3 is True)
            else:
                by[c]["false_alarm_when_correct"].append(q3 is True)
            by[c]["clarify"].append(yn(d.get("Q4_CLARIFY (Y/N)")) is True)
            try:
                by[c]["useful"].append(float(d.get("USEFUL (1-5)")))
            except Exception:  # noqa: BLE001
                pass
            t = minutes(d.get("START_TIME"), d.get("END_TIME"))
            if t is not None:
                by[c]["minutes"].append(t)
            by[c]["shown_wrong"].append(bool(k["shown_wrong"]))
    res = {"n_case_answers": n, "conditions": {}}
    for c, v in sorted(by.items()):
        res["conditions"][c] = {k_: ({"mean": round(sum(x) / len(x), 3), "n": len(x)} if x else None) for k_, x in v.items()}
    (SUB / "locked" / "practitioner_analysis.json").write_text(json.dumps(res, indent=1))
    print(json.dumps(res, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
