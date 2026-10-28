"""
Fresh-wording writing pack v4 (Improvement_Plan WS4.4): an independent annotator who has never seen the checker
writes sentences in the constructions where ReCheck failed: related fields, negation, units, temporal scope,
undocumented meanings and framing prose. Same four versions as v3 (original, paraphrase, corruption, provocation).

Borrowers: development-role borrowers never used by M1, practice, mechanism v1-v3 or the context packs.

    python build_wording_pack_v4.py --per-dataset 15
Output: Submission/packs/Wording_writing_pack_v4.zip (to the annotator)
        Submission/mechanism/annotation_key_wording_v4.jsonl (LOCAL ONLY)
"""
from __future__ import annotations

import argparse
import json
import random
import re
import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_heldout_writing_pack import README as V3_README  # noqa: E402
from build_mechanism_pack import OUT as MECH, _attrs, used_borrowers  # noqa: E402
from build_pack_m1 import SUB, assert_blind  # noqa: E402
from evidence import EvidenceType  # noqa: E402
from evidence_builder import render_evidence_v2  # noqa: E402

PACKS = SUB / "packs"
SEED = 20261013
RELS = {"taiwan": ["RELATED_FIELD", "NEGATION", "UNIT", "TEMPORAL_SCOPE", "UNDOCUMENTED_MEANING", "FRAMING"],
        "home_credit": ["RELATED_FIELD", "NEGATION", "UNIT", "UNDOCUMENTED_MEANING", "FRAMING"]}
REL_DEF = {
    "RELATED_FIELD": "the recorded value of the TARGET field, naming that field correctly (the evidence contains "
                     "related fields such as other amounts, bills, payments or limits that must not be confused with it)",
    "NEGATION": "whether the TARGET driver raises or lowers the model's risk score, written WITH a negation "
                "(e.g. 'does not ...', 'is not ...', 'no ...')",
    "UNIT": "the TARGET field's value with its unit, or converted correctly into another unit (e.g. days into years)",
    "TEMPORAL_SCOPE": "which statements or months the TARGET series covers, or what it shows over exactly those months",
    "UNDOCUMENTED_MEANING": "the TARGET item, whose meaning, code label, direction or currency the source does NOT document",
    "FRAMING": "a general framing sentence for a credit officer (how to read the model's summary) that states "
               "nothing about this borrower's data",
}
ERR_HINT = {
    "RELATED_FIELD": "the right number but the name of a RELATED, different field (e.g. a goods price called the "
                     "credit amount, a bill called a payment, a limit called a balance)",
    "NEGATION": "a negation that makes the direction wrong (e.g. 'does not lower' for a driver that lowers the score)",
    "UNIT": "a wrong unit or scale (e.g. thousands for units, days read as years, a ratio given as a percentage of "
            "something else)",
    "TEMPORAL_SCOPE": "a month or period that the evidence does not cover, or a change over months not shown",
    "UNDOCUMENTED_MEANING": "a definite meaning, label, direction or currency that the source does not document",
    "FRAMING": "framing prose that slips in ONE unsupported claim about this borrower (e.g. that they are reliable)",
}
PROV_HINT = {
    "RELATED_FIELD": "a correct sentence that also mentions a related field in order to DISTINGUISH it",
    "NEGATION": "a correct double negation or contrast (e.g. 'is not a factor that lowers risk; it raises it')",
    "UNIT": "a correct but unusual conversion or rounding (e.g. about 12.6 years for -4,610 days)",
    "TEMPORAL_SCOPE": "a correct statement about the months shown that uses words like 'over time' or 'recently'",
    "UNDOCUMENTED_MEANING": "a sentence that mentions a possible meaning only to say it is not documented",
    "FRAMING": "correct framing prose that uses risk words ('risk', 'default', 'raises') without any claim about "
               "this borrower",
}
README = V3_README.replace("# Held-out sentence-writing task", "# Wording sentence-writing task (round 4)") \
    .replace("`Heldout_writing_workbook.xlsx`", "`Wording_writing_workbook.xlsx`") \
    .replace("Estimated time: 5-6 hours for 240 sentences.", "This round focuses on six kinds of wording: "
             "related fields, negation, units, the months a series covers, undocumented meanings and general "
             "framing sentences. For FRAMING rows there is no target item.")


def _num(it):
    return isinstance(it.value, (int, float)) and not isinstance(it.value, bool)


def target(eo, rel, rng):
    facts = eo.by_type(EvidenceType.OBSERVED_FACT)
    if rel == "RELATED_FIELD":
        c = [f for f in facts if _num(f) and re.search(r"AMT|BILL|PAY_AMT|LIMIT|INCOME|CREDIT|ANNUITY|GOODS",
                                                         str(f.source_field), re.I)]
    elif rel == "NEGATION":
        c = [a for a in _attrs(eo) if abs(a.contribution or 0) >= 0.05]
    elif rel == "UNIT":
        c = [f for f in facts if _num(f) and re.search(r"dollar|currency|days|years|month", str(f.unit), re.I)]
    elif rel == "TEMPORAL_SCOPE":
        series = {}
        for f in facts:
            if f.series and f.period_order is not None:
                series.setdefault(f.series, []).append(f)
        c = [sorted(v, key=lambda x: -x.period_order) for v in series.values() if len(v) >= 2]
        return [x.eid for x in rng.choice(c)] if c else None
    elif rel == "UNDOCUMENTED_MEANING":
        c = [f for f in facts if re.search(r"not (?:documented|defined)", f"{f.unit} {f.value_label} {f.observed_at}", re.I)]
    else:
        return []
    return [rng.choice(sorted(c, key=lambda x: x.eid)).eid] if c else None


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from build_pack_m1 import extra_borrowers
    ap = argparse.ArgumentParser()
    ap.add_argument("--per-dataset", type=int, default=15)
    a = ap.parse_args()
    keyp = MECH / "annotation_key_wording_v4.jsonl"
    if keyp.exists():
        print(f"{keyp} exists: built once")
        return 1
    rng = random.Random(SEED)
    exclude = used_borrowers()
    for p in list(MECH.glob("mechanism_pack_v*.jsonl")) + list(MECH.glob("context_pack_v*.jsonl")):
        exclude |= {json.loads(l).get("borrower_id") for l in open(p)}
    v3 = MECH / "annotation_key_mechanism_v3.jsonl"
    if v3.exists():
        exclude |= {json.loads(l)["borrower_id"] for l in open(v3)}
    rows, key, evid, k = [], [], [], 0
    for ds in ("taiwan", "home_credit"):
        rels = RELS[ds]
        made = 0
        for eo, _ in extra_borrowers(ds, exclude, a.per_dataset * 3, rng):
            if made >= a.per_dataset:
                break
            got = []
            for _ in range(len(rels)):
                rel = rels[k % len(rels)]
                k += 1
                t = target(eo, rel, rng)
                if t is not None and rel not in [g[0] for g in got]:
                    got.append((rel, t))
                if len(got) == 2:
                    break
            if len(got) < 2:
                continue
            made += 1
            bcode = f"W{len(evid) + 1:02d}"
            evid.append((bcode, ds, render_evidence_v2(eo)))
            for j, (rel, ids) in enumerate(got):
                base = f"WW-{bcode}-{'AB'[j]}"
                for ver in ("original", "paraphrase", "corruption", "provocation"):
                    what = {"original": f"A correct sentence about {REL_DEF[rel]}.",
                            "paraphrase": "The same correct content as your 'original', in clearly different words.",
                            "corruption": f"Exactly one error: {ERR_HINT[rel]}.",
                            "provocation": f"A CORRECT sentence that may look wrong, e.g. {PROV_HINT[rel]}."}[ver]
                    rows.append([f"{base}-{ver[:4]}", bcode, ds.replace("_", " "), rel, ", ".join(ids) or "-", ver,
                                 what, "", "", ""])
            key.append({"borrower_code": bcode, "borrower_id": eo.borrower_id, "dataset": ds,
                        "relations": [{"relation": r, "target_ids": i} for r, i in got],
                        "evidence": eo.to_dict(include_groups=True)})
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    for line in README.splitlines():
        ws0.append([line])
    ws0.column_dimensions["A"].width = 120
    we = wb.create_sheet("Evidence")
    we.append(["borrower_code", "dataset", "evidence"])
    for b, ds, ev in evid:
        we.append([b, ds.replace("_", " "), ev])
    we.column_dimensions["C"].width = 140
    ww = wb.create_sheet("Write")
    ww.append(["annotator_code", "", "date", "", "used_AI_tool (must be N)", ""])
    ww.append([])
    ww.append(["item_code", "borrower_code", "dataset", "relation", "target_evidence", "version", "what_to_write",
               "TEXT", "INTENDED_ERROR (corruption only)", "INVITED_OBJECTION (provocation only)"])
    for r in rows:
        ww.append(r)
    bold, head, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    for cell in ww[3]:
        cell.font, cell.fill = bold, head
    for ws in (we, ww):
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = wrap
    for col, w in zip("ABCDEFGHIJ", (16, 8, 12, 18, 14, 12, 50, 60, 30, 30)):
        ww.column_dimensions[col].width = w
    for row in ww.iter_rows(min_row=4, min_col=8, max_col=10):
        for cell in row:
            cell.fill = fill
    ww.freeze_panes = "H4"
    bids = [x["borrower_id"] for x in key]
    for ws in (we, ww):
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert_blind(cell.value, bids)
    stage = PACKS / "Wording_writing_pack_v4"
    shutil.rmtree(stage, ignore_errors=True)
    stage.mkdir(parents=True)
    wb.save(stage / "Wording_writing_workbook.xlsx")
    (stage / "README_writing_task.md").write_text(README)
    z = PACKS / "Wording_writing_pack_v4.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(stage.iterdir()):
            zf.write(f, f"Wording_writing_pack_v4/{f.name}")
    with open(keyp, "w") as f:
        for x in key:
            f.write(json.dumps(x, default=str) + "\n")
    from collections import Counter
    print(z, len(evid), "borrowers,", len(rows), "sentences", dict(Counter(r[3] for r in rows if r[5] == "original")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
