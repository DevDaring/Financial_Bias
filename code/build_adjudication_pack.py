"""
Author adjudication pack for M1 (Future_Plan §B P0 / Research_Plan §11 days 1-3).

Two sheets to decide, in one workbook:
  A_Disagreed_66     primary natural items where the two raters disagreed (or one was unresolved)
  B_Contract_7       items BOTH raters called correct but the credit-evidence contract calls wrong
                     (code contradiction, or >= 2 LLM reviewers whose objection survived execution)

Blinding: no borrower id, no producer model, no judge model names, no group attribute.
Rater codes (AS/NK/SP) are kept so per-rater accuracy can be computed afterwards.

    .venv/bin/python build_adjudication_pack.py
Output: Submission/packs/M1_adjudication_pack.zip
        (M1_adjudication_workbook.xlsx, README_adjudication.md, contract_rules.md)
Import the returned workbook with import_adjudication.py.
"""
from __future__ import annotations

import csv
import json
import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import contract as C  # noqa: E402
from build_pack_m1 import assert_blind, blind, eo_from  # noqa: E402
from evidence_builder import render_evidence_v2  # noqa: E402
from score_m1 import M1, RATERS, SUB, find, read_sheet  # noqa: E402

OUT = SUB / "packs"
DEC = ["Y", "N", "UNRESOLVED"]
TYPES = ["invented_factor", "sign_reversal", "numeric_or_unit_error", "unsupported_trend",
         "unsupported_causal_or_action", "omission_of_essential_driver", "other_material_error"]

README = """# M1 adjudication pack (author only)

You decide the reference label for **73 items**. Your decision becomes the development reference
truth for every verifier result, so decide from the EVIDENCE and the written rules only. Do not
use any LLM or tool output beyond what is shown here.

## Files
- `M1_adjudication_workbook.xlsx` - the two sheets you fill in, plus `Instructions`.
- `contract_rules.md` - the 12 contract rules and the annotation rubric's error types.

## Sheet A_Disagreed_66 (66 items)
The two raters gave different answers to "is there ANY material error?" (or one marked it
UNRESOLVED). Both raters' labels, quotes and notes are shown.
Fill in:
- `ADJUDICATED_ANY` - Y (at least one material error), N (no material error) or UNRESOLVED
  (the evidence or the rubric genuinely cannot settle it; explain why).
- `adjudicated_types` - if Y, one or more error types (see the drop-down list in Instructions),
  separated by `;`.
- `adjudication_note` - one sentence: the rule you applied and the decisive words.

## Sheet B_Contract_7 (7 items)
Both raters said "no material error", but the contract says the text breaks a rule. The
`contract_finding` column shows the rule and the quoted words. These are the cases that decide
whether the contract is stricter than the rubric. Typical patterns:
- a direction ("lowers the risk score") given to a field that is only an INPUT of a derived driver
  (e.g. the credit amount, when the attribution belongs to the credit-to-income ratio);
- a meaning given to an undocumented code ("code 0 ... no delay");
- one field's value attached to another field's name.
Fill in `AUTHOR_DECISION` (error / correct / unresolved) and `rule_note`. If you decide
"correct", write the rule you are applying (e.g. "a direction for an input of a derived driver is
acceptable when the driver is named in the same sentence"), because the same rule must then be
applied to every item and to the contract.

## Rules for the adjudicator
1. Decide each item independently; do not look at verifier outputs or other files.
2. A rounding difference, a documented unit conversion or a disclaimer that denies causation is
   NOT an error.
3. Write a rule note whenever your decision depends on an interpretation; keep the notes consistent.
4. Return the workbook unchanged in structure (do not add, delete or reorder rows or columns).
Estimated time: 3-4 hours.
"""


def rules_md() -> str:
    lines = ["# Contract rules (contract/1.1) and rubric error types", "",
             "| Rule | Forbids | Rubric error type |", "|---|---|---|"]
    for k, (txt, et) in C.RULES.items():
        lines.append(f"| {k} | {txt} | {et} |")
    lines += ["", "Rubric error types: " + ", ".join(TYPES) + ".",
              "", "A material error is one that would mislead a credit officer about the model's evidence: "
                  "a wrong value, direction, ranking, horizon, scale, code meaning, an invented factor, an "
                  "unsupported trend or causal claim, or the omission of the probability or strongest driver."]
    return "\n".join(lines) + "\n"


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    OUT.mkdir(parents=True, exist_ok=True)
    adj = list(csv.DictReader(open(M1 / "adjudication_m1.csv")))
    conf = list(csv.DictReader(open(M1 / "contract_label_conflicts_for_adjudication.csv")))
    replay = json.load(open(M1 / "replay_objections_dev.json"))
    cmap = {c["item_id"]: c for c in replay["contract_vs_label_conflicts"]}
    key = {r["item_id"]: r for r in map(json.loads, open(SUB / "annotation_key_m1.jsonl")) if not r.get("_meta")}
    latest = {}
    for line in open(M1 / "records.jsonl"):
        r = json.loads(line)
        latest[(r["borrower_id"], r["arm"], r["family"])] = r
    lab = {}
    for rt in RATERS:
        for i, row in read_sheet(find(SUB / "SAARTHI_Annotated", f"M1_annotation_sheet_{rt}"))["rows"].items():
            lab.setdefault(i, {})[rt] = row
    bids = sorted({k["borrower_id"] for k in key.values()})

    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    for line in README.splitlines():
        ins.append([line])
    ins.append([])
    ins.append(["Error types (for adjudicated_types):"])
    for t in TYPES:
        ins.append([t])
    ins.column_dimensions["A"].width = 120
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    fill_in = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    # ---- sheet A
    wa = wb.create_sheet("A_Disagreed_66")
    cols = ["item_id", "dataset", "explanation", "evidence", "rater1", "r1_ANY", "r1_types", "r1_quote", "r1_notes",
            "rater2", "r2_ANY", "r2_types", "r2_quote", "r2_notes", "ADJUDICATED_ANY", "adjudicated_types",
            "adjudication_note"]
    wa.append(cols)
    for r in adj:
        k = key.get(r["item_id"], {})
        eo = None
        rec = latest.get((k.get("borrower_id"), k.get("arm"), k.get("family")))
        if rec:
            eo = eo_from(rec["evidence"])
        expl = blind(r["explanation"], eo) if eo else r["explanation"]
        wa.append([r["item_id"], r["dataset"], expl, r["evidence"], r["rater1"], r["r1_ANY"], r["r1_types"],
                   r["r1_quote"], r["r1_notes"], r["rater2"], r["r2_ANY"], r["r2_types"], r["r2_quote"],
                   r["r2_notes"], "", "", ""])
    # ---- sheet B
    wbb = wb.create_sheet("B_Contract_7")
    colsb = ["item_id", "dataset", "explanation", "evidence", "rater_labels", "rater_notes", "contract_finding",
             "AUTHOR_DECISION", "rule_note"]
    wbb.append(colsb)
    for c in conf:
        i = c["item_id"]
        k = key[i]
        rec = latest[(k["borrower_id"], k["arm"], k["family"])]
        eo = eo_from(rec["evidence"])
        rows = lab.get(i, {})
        rl = "; ".join(f"{rt}: ANY={rows[rt].get('ANY_MATERIAL_ERROR')}" for rt in k["raters"] if rt in rows)
        rn = " | ".join(f"{rt}: {rows[rt].get('notes') or ''}" for rt in k["raters"] if rt in rows)
        cm = cmap.get(i, {})
        finding = []
        if c["code_rejects"] == "True":
            finding.append("Automatic check: a number written next to a field's name does not match that field's "
                           "value (it matches a different field).")
        seen = set()
        for obs in cm.get("objections", {}).values():
            for o in obs:
                q = (o.get("quote") or "")[:200]
                if q and q not in seen:
                    seen.add(q)
                    finding.append(f"{o['rule']}: \"{blind(q, eo)}\"")
        wbb.append([i, rec["dataset"], blind(W_clean(rec["text"]), eo), render_evidence_v2(eo), rl, rn,
                    "\n".join(finding[:4]), "", ""])
    for ws, n_dec in ((wa, 3), (wbb, 2)):
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
        widths = {"explanation": 70, "evidence": 80, "contract_finding": 60}
        for j, name in enumerate([c.value for c in ws[1]], start=1):
            letter = ws.cell(row=1, column=j).column_letter
            ws.column_dimensions[letter].width = widths.get(name, 16 if j <= len(cols) - n_dec else 28)
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                for cell in row:
                    cell.alignment = wrap
                    if name in ("ADJUDICATED_ANY", "adjudicated_types", "adjudication_note", "AUTHOR_DECISION",
                                "rule_note"):
                        cell.fill = fill_in
        ws.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"Y,N,UNRESOLVED"', allow_blank=True)
    wa.add_data_validation(dv)
    dv.add(f"O2:O{wa.max_row}")
    dv2 = DataValidation(type="list", formula1='"error,correct,unresolved"', allow_blank=True)
    wbb.add_data_validation(dv2)
    dv2.add(f"H2:H{wbb.max_row}")
    for ws in (wa, wbb):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    assert_blind(cell.value, bids)

    stage = OUT / "M1_adjudication_pack"
    shutil.rmtree(stage, ignore_errors=True)
    stage.mkdir(parents=True)
    wb.save(stage / "M1_adjudication_workbook.xlsx")
    (stage / "README_adjudication.md").write_text(README)
    (stage / "contract_rules.md").write_text(rules_md())
    z = OUT / "M1_adjudication_pack.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(stage.iterdir()):
            zf.write(f, f"M1_adjudication_pack/{f.name}")
    print(f"{z}  ({len(adj)} disagreed + {len(conf)} contract conflicts)")
    return 0


def W_clean(t: str) -> str:
    import witness as W
    return W.clean(t)


if __name__ == "__main__":
    raise SystemExit(main())
