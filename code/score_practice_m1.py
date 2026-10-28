"""
Practice gate for the milestone-1 annotation (plan §1.4 lesson: a rater's
labels are counted only after a qualification check).

    python score_practice_m1.py path/to/M1_PRACTICE_AS.xlsx [more sheets ...]

Reads ANY_MATERIAL_ERROR from each returned practice sheet, compares it with
Submission/practice_key_m1.json (kept local), prints the score and whether the
rater passes, and writes the feedback text to send back (answers included).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

SUB = Path(__file__).resolve().parents[2] / "Submission"


def score(path: Path, key: dict) -> dict:
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Annotate"]
    head = [c.value for c in ws[1]]
    ci, ca = head.index("item_id"), head.index("ANY_MATERIAL_ERROR")
    got = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[ci]:
            got[str(row[ci]).strip()] = str(row[ca] or "").strip().upper()[:1]
    res, ok = [], 0
    for pid, k in sorted(key["items"].items()):
        g = got.get(pid, "")
        hit = g == k["ANY"]
        ok += hit
        res.append((pid, g or "-", k["ANY"], hit, k["answer"]))
    blank = sum(1 for r in res if r[1] == "-")
    return {"score": ok, "n": len(res), "blank": blank,
            "passed": ok >= key["pass_mark"] and blank == 0, "rows": res}


def main() -> int:
    key = json.load(open(SUB / "practice_key_m1.json"))
    for p in map(Path, sys.argv[1:]):
        r = score(p, key)
        print(f"{p.name}: {r['score']}/{r['n']} (blank {r['blank']}) -> "
              f"{'PASS' if r['passed'] else 'NOT PASSED'} (pass mark {key['pass_mark']})")
        fb = [f"Practice feedback for {p.stem}: {r['score']}/{r['n']} correct.\n"]
        for pid, g, t, hit, ans in r["rows"]:
            fb.append(f"{pid}: you {g}, answer {t} {'✓' if hit else '✗'} — {ans}")
        out = p.with_name(p.stem + "_feedback.txt")
        out.write_text("\n".join(fb))
        print(f"  feedback -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
