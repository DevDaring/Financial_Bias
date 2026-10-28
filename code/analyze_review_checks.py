"""
Human review checks (Improvement_Plan WS3, WS1.5). Reads ONLY label/answer columns; annotator, checker and participant
tabs (identity, dates, time, role, experience) are never opened.

  1. independent rating of locked-v1 items vs the earlier final labels; locked-v1 primary recomputed with them
  2. context check: human truth judgement vs the construction label of each focal sentence
  3. reader parse check: accuracy of the reader's binding, operation, asserted value and polarity

    python analyze_review_checks.py -> Submission/locked/review_checks.json
"""
from __future__ import annotations

import json
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dev_sets import locked_items  # noqa: E402
from score_m1 import SUB, yn  # noqa: E402

D = SUB / "packs" / "Annotated" / "IDBI_Annotator_Checks" / "Returned"


def sheet(path, name, header_row=0):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[name].iter_rows(values_only=True))
    head = [str(c) for c in rows[header_row]]
    return [dict(zip(head, r)) for r in rows[header_row + 1:] if r and r[0]]


def main() -> int:
    res = {}
    # 1. independent rating
    key = {json.loads(l)["item_id"]: json.loads(l) for l in open(SUB / "annotation_key_independent_v1.jsonl")}
    ir = {r["item_id"]: yn(r.get("ANY_MATERIAL_ERROR")) for r in sheet(next(D.rglob("INDEPENDENT_rating_sheet.xlsx")), "Annotate")}
    final = {x["item_id"]: x["truth"] for x in locked_items()}
    tmpl_final = False                     # all locked-v1 templates were labelled correct
    agg = defaultdict(Counter)
    ir_by_src = {}
    for sid, k in key.items():
        v = ir.get(sid)
        if v is None:
            continue
        f = final.get(k["source_item_id"], tmpl_final) if k["arm"] != "template" else tmpl_final
        agg[k["reason"]]["agree" if v == f else "differ"] += 1
        agg[k["reason"]][f"ir_{'error' if v else 'correct'}_final_{'error' if f else 'correct'}"] += 1
        if k["arm"] != "template":
            ir_by_src[k["source_item_id"]] = v
    res["independent"] = {k: dict(v) for k, v in agg.items()}
    # locked-v1 primary with independent labels replacing disagreements and revised items
    truth = dict(final)
    for sid, k in key.items():
        if k["reason"] in ("disagreement", "revised_after_feedback") and k["source_item_id"] in ir_by_src:
            truth[k["source_item_id"]] = ir_by_src[k["source_item_id"]]
    items = {x["item_id"]: x for x in locked_items()}
    strata = defaultdict(list)
    for i, x in items.items():
        strata[x["rec"]["stratum"]].append(i)
    rng = random.Random(20261020)
    draws = [[i for v in strata.values() for i in (rng.choice(v) for _ in v)] for _ in range(3000)]
    runs = SUB / "locked" / "fresh_narratives_v1_runs"
    out = {"n_changed": sum(truth[i] != final[i] for i in truth), "n_wrong": sum(truth.values())}
    for j in ("qwen3_next", "gemma27_vllm"):
        a = {json.loads(l)["item_id"]: json.loads(l).get("accepted") for l in open(runs / f"exec_R__{j}.jsonl")}
        b = {json.loads(l)["item_id"]: json.loads(l).get("accepted") for l in open(runs / f"facts_judge__{j}.jsonl")}
        o = {}
        for kind, w in (("wrong_release", True), ("correct_release", False)):
            def rate(its, d):
                x = [i for i in its if truth[i] == w]
                return sum(d.get(i) is True for i in x) / max(1, len(x))
            pt = rate(list(truth), a) - rate(list(truth), b)
            bs = sorted(rate(s, a) - rate(s, b) for s in draws)
            o[kind] = {"diff": round(pt, 4), "ci95": [round(bs[75], 4), round(bs[2924], 4)]}
        out[j] = o
    res["locked_with_independent_labels"] = out
    # 2. context check
    ck = {json.loads(l)["item_id"]: json.loads(l) for l in open(SUB / "mechanism" / "annotation_key_context_check.jsonl")}
    cc = defaultdict(Counter)
    flagged = []
    for r in sheet(next(D.rglob("CONTEXT_check_sheet.xlsx")), "Check"):
        k = ck.get(r["item_id"])
        if not k:
            continue
        human_correct, unclear = yn(r.get("SENTENCE_CORRECT")), yn(r.get("UNCLEAR"))
        intended_correct = not k["intended_error"]
        tag = "unclear" if unclear else ("match" if human_correct == intended_correct else "mismatch")
        cc[k["relation"]][tag] += 1
        cc["ALL"][tag] += 1
        if tag != "match":
            flagged.append({"pack": k["pack"], "base_id": k["base_id"], "context_item_id": k["context_item_id"], "tag": tag})
    res["context_check"] = {k: dict(v) for k, v in cc.items()}
    res["context_flagged"] = flagged
    # 3. reader parse check
    pk = {json.loads(l)["item_id"]: json.loads(l) for l in open(SUB / "annotation_key_parse_check.jsonl")}
    pc = defaultdict(Counter)
    for r in sheet(next(D.rglob("READER_parse_check_sheet.xlsx")), "Check"):
        k = pk.get(r["item_id"])
        if not k:
            continue
        for f in ("BINDING_OK", "OPERATION_OK", "ASSERTED_OK", "POLARITY_OK"):
            v = yn(r.get(f))
            if v is not None:
                pc[f]["ok" if v else "wrong"] += 1
                pc[f"{f}|{k['operation']}"]["ok" if v else "wrong"] += 1
        allok = all(yn(r.get(f)) for f in ("BINDING_OK", "OPERATION_OK", "ASSERTED_OK", "POLARITY_OK"))
        pc["ALL_FOUR_OK"]["ok" if allok else "wrong"] += 1
    res["parse_check"] = {k: dict(v) for k, v in pc.items()}
    (SUB / "locked" / "review_checks.json").write_text(json.dumps(res, indent=1))
    print(json.dumps({k: v for k, v in res.items() if k != "context_flagged"}, indent=1)[:4000])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
