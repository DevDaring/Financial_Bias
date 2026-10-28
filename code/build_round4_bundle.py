"""
Round-4 human tasks (response to the review), packed as ONE zip with one folder per task.

Task 1  Independent adjudication of the fresh-test disagreements (the protocol named in the prespecification:
        an adjudicator who is neither an author nor one of the four raters, blind to every verifier output).
        Each row shows the evidence, the explanation and the two earlier ratings as "Rater A" / "Rater B"
        (material error, error types, quoted passage). Rater identities, rater notes and the provisional
        decision are NOT shown.
Task 2  Meaning check of rewritten sentences (paraphrase stress test): does the rewrite say exactly what the
        original says? A stratified random sample over rewriter x style.

    python build_round4_bundle.py
Outputs: Submission/packs/Annotation_round4_bundle.zip (send)
         Submission/annotation_key_round4.jsonl (KEEP LOCAL)
"""
from __future__ import annotations

import csv
import json
import random
import shutil
import sys
import zipfile
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_pack_m1 as PM  # noqa: E402
import build_fresh_pack_v2 as FP  # noqa: E402
import feature_semantics as FS  # noqa: E402
from analyze_locked import TYPES  # noqa: E402
from build_pack_m1 import SUB, assert_blind, blind, eo_from  # noqa: E402
from evidence_builder import render_evidence_v2  # noqa: E402
from score_m1 import read_sheet, yn  # noqa: E402

SEED = 20261013
PACKS = SUB / "packs"
LABELS = PACKS / "Annotated" / "IDBI_Annotator_Checks"
N_PARA = 120

README_FIRST = """Annotation round 4 (two tasks)
==============================

This zip has one folder per task. Each folder has its own README, a sheet to fill (yellow columns) and the
reference documents. The two tasks are independent; one person may do both.

  Task1_Independent_adjudication/   {n1} items, about 2-3 hours
      Who: a person who is NOT an author of the study and who did NOT rate these explanations before.
  Task2_Paraphrase_check/           {n2} sentence pairs, about 45 minutes
      Who: a person who is NOT an author of the study.

Rules for both tasks
  - Work alone. Do not discuss items with the authors or with the earlier raters before returning the sheets.
  - Do not use ChatGPT or any other AI tool.
  - The material is for this study only; do not copy or share it.

Return
  Task1_Independent_adjudication/ADJUDICATION_sheet.xlsx
  Task2_Paraphrase_check/PARAPHRASE_check_sheet.xlsx
Only the yellow columns are used. Please do not rename the files, add rows or reorder rows.
"""

ADJ_README = """Task 1: independent adjudication
================================

Two trained raters labelled each of these credit-risk explanations against the model evidence and DISAGREED on
whether it contains a material error. You decide. You are the final, independent judge: you are not bound by
either rater, and you have not seen any software verdict (none is shown on purpose).

For each row
  1. Read the evidence (left) and the explanation.
  2. Read the two ratings (grey columns). Rater A and Rater B are not the same two people on every row.
  3. Apply ANNOTATION_GUIDELINES.md (the same rules the raters used; read the addenda at the end) and
     FEATURE_GLOSSARY.md (what each field means).
  4. Fill the yellow columns:
       FINAL_MATERIAL_ERROR   Y if the explanation contains at least one material error, N if not
       FINAL_ERROR_TYPE       if Y: invented factor | sign reversal | number or unit | unsupported trend |
                              unsupported cause or action | omits the essential driver | other
       ERROR_QUOTE            if Y: the words that are wrong (copy them)
       CANNOT_DECIDE          Y only if the evidence truly does not allow a decision
       NOTES                  optional, one line

A material error is one that would mislead a credit officer about the model's evidence: a wrong value, a wrong
direction, the wrong strongest driver, a probability or contribution stated on the wrong scale, a number given to
the wrong field, a meaning given to a code the source does not document, or a factor that is not in the evidence.
Plain or everyday wording is not an error. Contributions are on the raw log-odds scale, not percentages.
"""

PARA_README = """Task 2: meaning check of rewritten sentences
============================================

Each row shows an ORIGINAL sentence about a borrower's credit-risk evidence and a REWRITE of it produced by a
language model. We need to know whether the rewrite says exactly the same thing. You do not need the evidence:
compare the two sentences only. The original may itself be right or wrong; that does not matter here.

Fill the yellow columns
  SAME_MEANING   Y if the rewrite states the same facts, numbers, directions (raises / lowers), negations,
                 fields, time periods and units as the original; N if anything that matters changed
  WHAT_CHANGED   if N: number | direction | negation | field or name | time period | unit or scale |
                 added information | removed information | other
  NOTES          optional

Different words, word order, voice or everyday names for the same field are fine (Y), as long as the meaning is
the same. A hedge ("about", "roughly") added to an exact number counts as a change only if it changes what a
credit officer would understand.
"""


def task1(stage: Path, rng):
    key = {k["item_id"]: k for k in (json.loads(l) for l in open(SUB / "annotation_key_fresh_v2.jsonl")) if "item_id" in k}
    disagreements = [r["item_id"] for r in csv.DictReader(open(SUB / "locked" / "adjudication_fresh_v2.csv"))]
    sheets = {}
    for q in FP.BL.RATERS:
        hits = sorted(LABELS.rglob(f"FRESH_annotation_sheet_{q}*.xlsx"))
        sheets[q] = read_sheet(hits[-1])["rows"]       # label columns only; the Annotator tab is never read
    recs = {json.loads(l)["item_id"]: json.loads(l) for l in open(SUB / "locked" / "fresh_natural_v2_final.jsonl")}
    rows, keyrows = [], []
    items = list(disagreements)
    rng.shuffle(items)

    def rating(r):
        types = [t.replace("_", " ") for t in TYPES if yn(r.get(t))]
        return ("Y" if yn(r.get("ANY_MATERIAL_ERROR")) else "N" if yn(r.get("ANY_MATERIAL_ERROR")) is False else "?",
                "; ".join(types) or "-", str(r.get("error_quote") or "")[:300],
                "Y" if yn(r.get("UNRESOLVED")) else "")
    for n, i in enumerate(items, 1):
        k = key[i]
        rec = recs[k["source_item_id"]]
        eo = eo_from(rec["evidence"])
        a, b = [rating(sheets[q][i]) for q in k["raters"]]
        if rng.random() < 0.5:
            a, b = b, a
        sid = f"AJ-{n:03d}"
        rows.append([sid, PM.DATASET_LABEL[rec["dataset"]], render_evidence_v2(eo), blind(rec["text"], eo),
                     *a, *b, "", "", "", "", ""])
        keyrows.append({"task": 1, "item_id": sid, "fresh_item_id": i, "source_item_id": k["source_item_id"]})
    bids = {recs[key[i]["source_item_id"]]["borrower_id"] for i in items}
    for r in rows:
        assert_blind(" ".join(str(x) for x in r), bids)
    head = ["item_id", "dataset", "evidence", "explanation",
            "A_material_error", "A_error_types", "A_quote", "A_unresolved",
            "B_material_error", "B_error_types", "B_quote", "B_unresolved",
            "FINAL_MATERIAL_ERROR", "FINAL_ERROR_TYPE", "ERROR_QUOTE", "CANNOT_DECIDE", "NOTES"]
    widths = [9, 12, 70, 60, 9, 20, 28, 9, 9, 20, 28, 9, 12, 20, 28, 10, 28]
    save_sheet(stage / "ADJUDICATION_sheet.xlsx", "Adjudicate", head, rows, widths, first_yellow=12, grey=(4, 12),
               yn_cols=("FINAL_MATERIAL_ERROR", "CANNOT_DECIDE"))
    (stage / "README.txt").write_text(ADJ_README)
    (stage / "ANNOTATION_GUIDELINES.md").write_text(PM.GUIDE + FP.ADDENDUM)
    (stage / "FEATURE_GLOSSARY.md").write_text(FS.glossary_markdown())
    return keyrows


def task2(stage: Path, rng):
    pack = [json.loads(l) for l in open(SUB / "mechanism" / "paraphrase_pack.jsonl")]
    orig = {r["source_id"]: r for r in pack if r["style"] == "original"}
    cells = {}
    for r in pack:
        if r["style"] != "original":
            cells.setdefault((r["rewriter"], r["style"]), []).append(r)
    per = N_PARA // max(1, len(cells))
    pick = [x for c in sorted(cells) for x in rng.sample(cells[c], min(per, len(cells[c])))]
    rng.shuffle(pick)
    rows, keyrows = [], []
    for n, r in enumerate(pick, 1):
        sid = f"PC-{n:03d}"
        o = orig[r["source_id"]]
        eo = eo_from(r["evidence"])
        rows.append([sid, blind(o["text"], eo), blind(r["text"], eo), "", "", ""])
        keyrows.append({"task": 2, "item_id": sid, "paraphrase_item_id": r["item_id"], "rewriter": r["rewriter"],
                        "style": r["style"], "truth": r["truth"], "relation": r["relation"]})
    head = ["item_id", "ORIGINAL sentence", "REWRITE", "SAME_MEANING", "WHAT_CHANGED", "NOTES"]
    save_sheet(stage / "PARAPHRASE_check_sheet.xlsx", "Check", head, rows, [9, 60, 60, 12, 20, 28], first_yellow=3,
               grey=None, yn_cols=("SAME_MEANING",))
    (stage / "README.txt").write_text(PARA_README)
    return keyrows


def save_sheet(path, title, head, rows, widths, first_yellow, grey, yn_cols):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(head)
    for r in rows:
        ws.append(r)
    hdr, yel, gry = PatternFill("solid", fgColor="1F2937"), PatternFill("solid", fgColor="FEF3C7"), \
        PatternFill("solid", fgColor="E5E7EB")
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), hdr, Alignment(wrap_text=True)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for row in ws.iter_rows(min_row=2):
        for j, c in enumerate(row):
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if j >= first_yellow:
                c.fill = yel
            elif grey and grey[0] <= j < grey[1]:
                c.fill = gry
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    for nm in yn_cols:
        col = get_column_letter(head.index(nm) + 1)
        dv.add(f"{col}2:{col}{len(rows) + 1}")
    ws.freeze_panes = ws.cell(row=2, column=first_yellow + 1)
    meta = wb.create_sheet("Declaration")
    meta.append(["I worked alone and used no AI tool (Y/N)", ""])
    meta.column_dimensions["A"].width = 45
    wb.save(path)


def main() -> int:
    rng = random.Random(SEED)
    root = PACKS / "Annotation_round4"
    shutil.rmtree(root, ignore_errors=True)
    t1, t2 = root / "Task1_Independent_adjudication", root / "Task2_Paraphrase_check"
    t1.mkdir(parents=True)
    t2.mkdir(parents=True)
    k1 = task1(t1, rng)
    k2 = task2(t2, rng)
    (root / "README_FIRST.txt").write_text(README_FIRST.format(n1=len(k1), n2=len(k2)))
    z = PACKS / "Annotation_round4_bundle.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(root.rglob("*")):
            if f.is_file():
                zf.write(f, f"Annotation_round4/{f.relative_to(root)}")
    (SUB / "annotation_key_round4.jsonl").write_text("\n".join(json.dumps(x) for x in k1 + k2) + "\n")
    print(z, "task1", len(k1), "task2", len(k2), Counter((x["rewriter"], x["style"]) for x in k2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
