"""
Milestone-1 annotation pack (plan §5.1-5.3), built ONCE from Submission/m1.

Items
  natural    every successful M1 candidate: template + producers x {grounded, typed}
  challenge  validated single-error mutations of template text (known answer,
             6 types) + valid paraphrases of template text (known clean).
             Template bases are correct by construction from the evidence; the
             mutation sites are chosen and verified from the evidence object,
             never with the checker being evaluated.
  duplicate  ~8% hidden repeats per annotator (intra-rater consistency)

Design
  - each item is labelled by exactly two of the three annotators (AS, NK, SP),
    balanced over the three pairs within dataset x source strata;
  - each annotator gets their own shuffled sheet and a 10-item PRACTICE sheet
    without answers. Main labels count only if practice scores >= 8/10
    (score_practice_m1.py). Practice borrowers are development-role borrowers
    NOT used in the main pack;
  - new columns: UNRESOLVED (+ reason), reasonable_disagreement,
    error_evidence_ids, per plan §5.2 and Guerdan et al. (rating
    indeterminacy) / Wang et al. (rationale-level reference).

Outputs
  Submission/annotation_pack_m1_<AS|NK|SP>.zip   -> send one to each annotator
  Submission/annotation_key_m1.jsonl             -> KEEP LOCAL (unblinding)
  Submission/practice_key_m1.json                -> KEEP LOCAL (practice answers)
"""
from __future__ import annotations

import json
import random
import re
import shutil
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Optional

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_annotation_pack as v1  # noqa: E402
import feature_semantics as FS  # noqa: E402
from evidence import EvidenceItem, EvidenceObject, EvidenceType  # noqa: E402
from evidence_builder import _fmt_value, render_evidence_v2  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
SUB = ROOT / "Submission"
M1 = SUB / "m1"
SEED = 20261003
RATERS = ["AS", "NK", "SP"]
PAIRS = [("AS", "NK"), ("AS", "SP"), ("NK", "SP")]
DUP_RATE = 0.08
PASS_MARK = 8

LABEL_COLS = [
    "invented_factor", "sign_reversal", "numeric_or_unit_error",
    "unsupported_trend", "unsupported_causal_or_action",
    "omission_of_essential_driver", "other_material_error",
    "ANY_MATERIAL_ERROR", "UNRESOLVED", "unresolved_reason",
    "reasonable_disagreement", "informative_for_credit_officer", "confidence",
    "error_quote", "error_evidence_ids", "notes",
]
YN_COLS = LABEL_COLS[:9] + ["reasonable_disagreement", "informative_for_credit_officer"]
DATASET_LABEL = {"taiwan": "Taiwan credit card", "home_credit": "Home Credit"}


# ---------------------------------------------------------------------------
def eo_from(d: dict) -> EvidenceObject:
    eo = EvidenceObject(borrower_id=d["borrower_id"], dataset=d["dataset"],
                        prediction_time=d.get("prediction_time"),
                        group_attributes=d.get("group_attributes", {}),
                        schema_version=d.get("schema_version", "evidence/2.0"),
                        provenance=d.get("provenance", {}))
    for it in d["items"]:
        it = dict(it)
        it["etype"] = EvidenceType(it["etype"])
        eo.add(EvidenceItem(**it))
    return eo


def blind(text: str, eo: EvidenceObject) -> str:
    """Remove the borrower id: 'taiwan-20077' in any case/separator, and a bare
    number only when it is introduced as an id ('borrower 20077', '#20077',
    'ID 20077'), never a free-standing number (an amount could equal the id)."""
    num = eo.borrower_id.rsplit("-", 1)[-1]
    ds = re.escape(eo.dataset).replace("_", "[-_ ]?")
    out = re.sub(rf"(?:\b(?:the\s+)?borrower\s+)?\b{ds}[-_ #]?{re.escape(num)}\b", "this borrower", text, flags=re.I)
    out = re.sub(rf"\b(?:the\s+)?borrower\s*(?:id|no\.?|number)?\s*[#:]?\s*#?{re.escape(num)}\b", "this borrower",
                 out, flags=re.I)
    out = re.sub(rf"(?:\bid\s*[#:]?\s*|#){re.escape(num)}\b", "this borrower", out, flags=re.I)
    return out


def assert_blind(text: str, bids) -> None:
    for b in bids:
        ds, num = b.rsplit("-", 1)
        if re.search(rf"{re.escape(ds).replace('_', '[-_ ]?')}[-_ #]?{num}\b", text, flags=re.I):
            raise SystemExit(f"blinding failed: {b} visible in a sheet cell")


def _name(a) -> str:
    return (a.feature_group or a.source_field or "").replace(" (derived)", "")


def _sentences(t: str):
    return [s for s in re.split(r"(?<=[.!?])\s+(?=[A-Z])", t.strip()) if s]


def top_attr(eo):
    at = eo.by_type(EvidenceType.MODEL_ATTRIBUTION)
    return max(at, key=lambda a: abs(a.contribution or 0)) if at else None


# ---------------------------------------------------------------------------
# challenge set: single-error mutations of template CLAIMS with verified answers
# ---------------------------------------------------------------------------
def _driver_claims(claims, eo):
    """attribution eid -> index of the template claim that states it."""
    ids = {a.eid for a in eo.by_type(EvidenceType.MODEL_ATTRIBUTION)}
    out = {}
    for k, c in enumerate(claims):
        for e in c.get("evidence_ids") or []:
            if e in ids and e not in out and len(c.get("evidence_ids") or []) <= 2:
                out[e] = k
    return out


def _all_values(eo):
    vals = []
    for it in eo.items:
        for v in (it.value, it.calibrated_pd, it.contribution):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                vals.append(float(v))
    return vals


def _series_shown(eo, f) -> bool:
    """f, or any of its inputs, belongs to a series with >= 2 shown periods."""
    facts = {x.eid: x for x in eo.by_type(EvidenceType.OBSERVED_FACT)}
    periods = defaultdict(set)
    for x in facts.values():
        if x.series:
            periods[x.series].add(x.period_order)
    stack, seen = [f], set()
    while stack:
        x = stack.pop()
        if x is None or x.eid in seen:
            continue
        seen.add(x.eid)
        if x.series and len(periods[x.series]) >= 2:
            return True
        stack += [facts.get(c) for c in (x.components or [])]
    return False


# wording pools: MAIN pack frames and PRACTICE-only frames never overlap, so the
# practice answers cannot teach raters to recognise main-pack challenge items
FRAMES = {
    "main": {
        "insert": ["{nm} also {verb} the model's risk score.",
                   "The recorded {nm} further {verb} the estimate.",
                   "In addition, {nm} {verb} the model's score relative to its reference.",
                   "{nm} is another factor that {verb} the risk score."],
        "trend": ["The borrower's {nm} has been rising over recent months.",
                  "{nm} has been falling month after month.",
                  "Over the past months, {nm} has gone up steadily.",
                  "{nm} has been declining over time."],
        "causal_up": ["Lowering {nm} will reduce the borrower's risk of default.",
                      "If the borrower reduced {nm}, the chance of default would fall.",
                      "{nm} is a cause of default for this borrower.",
                      "Bringing {nm} down would prevent a default."],
        "causal_dn": ["Raising {nm} will reduce the borrower's risk of default.",
                      "If the borrower increased {nm}, the chance of default would fall.",
                      "{nm} protects this borrower from default.",
                      "Keeping {nm} high will prevent a default."],
    },
    "practice": {
        "insert": ["{nm} likewise {verb} the score."],
        "trend": ["{nm} keeps climbing from one statement to the next."],
        "causal_up": ["Cutting {nm} would stop this borrower from defaulting."],
        "causal_dn": ["Increasing {nm} would stop this borrower from defaulting."],
    },
}


def _numeric_ok(f) -> bool:
    """A fact whose value can be changed to a plausible wrong value."""
    nm = (f.internal_name or "").upper()
    if f.value_label or not isinstance(f.value, float) or f.value == 0:
        return False                         # codes, text, zeros
    if nm.startswith(("FLAG_", "REG_", "LIVE_")) or f.value in (0.0, 1.0):
        return False                         # flags
    return True


def _mutate_value(v: float, rng: random.Random, vals) -> Optional[float]:
    """Wrong value in the SAME precision and a plausible range."""
    for _ in range(12):
        if abs(v) >= 1 and float(v).is_integer():
            nv = float(round(v * rng.choice([1.6, 0.6, 2.2, 0.4])))
            if abs(v) < 120:                  # ages, counts, small amounts: shift
                nv = float(v + rng.choice([-12, -7, 7, 12]))
        elif abs(v) < 1:
            dec = max(2, len(_fmt_value(v).split(".")[-1]) if "." in _fmt_value(v) else 2)
            nv = round(min(0.99, max(0.01, v + rng.choice([-0.3, -0.2, 0.2, 0.3]))), dec)
        else:
            nv = round(v * rng.choice([1.6, 0.6, 2.2]), 2)
        if nv != v and nv > 0 if v > 0 else nv != v:
            if not any(abs(nv - x) <= max(0.005 * abs(x), 1e-6) for x in vals):
                return nv
    return None


def _trend_ok(f) -> bool:
    nm = (f.internal_name or "").upper()
    return (f.value is not None and not f.value_label and not isinstance(f.value, str)
            and not nm.startswith(("AGE", "DAYS_", "FLAG_", "OWN_CAR_AGE", "BUREAU_DAYS"))
            and "age" not in (f.meaning or "").lower().split())


def mutate(claims, eo: EvidenceObject, etype: str, rng: random.Random, frames: str = "main"):
    """Return (mutated claim list, note) or None. Exactly one claim changes (or,
    for invented/trend/causal, exactly one claim is added; for omission, exactly
    the strongest driver's claim is removed), and the change contains ONLY the
    intended error. Never uses the checker under study."""
    F = FRAMES[frames]
    claims = [dict(c) for c in claims]
    attrs = sorted(eo.by_type(EvidenceType.MODEL_ATTRIBUTION), key=lambda a: -abs(a.contribution or 0))
    facts = {f.eid: f for f in eo.by_type(EvidenceType.OBSERVED_FACT)}
    dc = _driver_claims(claims, eo)
    text = " ".join(c["text"] for c in claims)
    if not attrs or not dc:
        return None
    insert_at = max(1, len(claims) - 2)      # before the scale note and the disclaimer

    if etype == "invented_factor":
        shown = {(f.internal_name or "") for f in facts.values()} | {(a.internal_name or "") for a in attrs}
        specs = FS.specs(eo.dataset)
        # also exclude what is computable from, or an input of, a shown field
        derivable = {k for k, sp in specs.items() if sp.components and set(sp.components) <= shown}
        inputs = {c for k in shown if k in specs for c in specs[k].components}
        pool = [s for k, s in specs.items() if k not in shown | derivable | inputs
                and s.concept not in ("UNMAPPED_FEATURE", "REGION_AND_HOUSING", "APPLICATION_DETAIL",
                                      "SOCIAL_CIRCLE_RECORD")
                and not any(al.lower() in text.lower() for al in [s.display] + s.aliases if len(al) > 3)]
        if not pool:
            return None
        s = rng.choice(sorted(pool, key=lambda x: x.internal))
        nm = s.display.replace(" (derived)", "")
        add = {"text": rng.choice(F["insert"]).format(nm=nm, verb=rng.choice(["raises", "lowers"])),
               "evidence_ids": []}
        return claims[:insert_at] + [add] + claims[insert_at:], f"invented driver {nm} (not in evidence): \"{add['text']}\""

    if etype == "sign_reversal":
        for a in attrs[1:] + attrs[:1]:
            k = dc.get(a.eid)
            if k is None:
                continue
            t = claims[k]["text"]
            if t.count(" raises the model's risk score") + t.count(" lowers the model's risk score") != 1:
                continue
            claims[k]["text"] = (t.replace(" raises the model's", " lowers the model's") if " raises the model's" in t
                                 else t.replace(" lowers the model's", " raises the model's"))
            return claims, f"direction of {_name(a)} reversed ({a.eid}): \"{claims[k]['text']}\""
        return None

    if etype == "numeric_or_unit_error":
        vals = _all_values(eo)
        for a in attrs:
            k = dc.get(a.eid)
            f = facts.get((a.components or [None])[0])
            if k is None or f is None or not _numeric_ok(f):
                continue
            fmt = _fmt_value(f.value)
            old = f"(recorded value {fmt} "
            t = claims[k]["text"]
            if t.count(old) != 1:
                continue
            nv = _mutate_value(f.value, rng, vals)
            if nv is None:
                continue
            nf = _fmt_value(float(nv))
            claims[k]["text"] = t.replace(old, f"(recorded value {nf} ", 1)
            return claims, f"value of {_name(a)} changed from {fmt} to {nf} ({f.eid})"
        return None

    if etype == "unsupported_trend":
        for a in attrs:
            f = facts.get((a.components or [None])[0])
            if f is None or _series_shown(eo, f) or not _trend_ok(f):
                continue                    # comparisons, codes, text, flags and ages are not used
            add = {"text": rng.choice(F["trend"]).format(nm=_name(a)), "evidence_ids": []}
            return claims[:insert_at] + [add] + claims[insert_at:], \
                f"trend claim on {_name(a)}: \"{add['text']}\""
        return None

    if etype == "unsupported_causal_or_action":
        for a in attrs:
            f = facts.get((a.components or [None])[0])
            if f is None or not _trend_ok(f):
                continue                    # no 'raising' a code, a flag or an age
            # wording follows the evidence direction: the ONLY error is causal
            pool = F["causal_up"] if a.direction == "increases_risk" else F["causal_dn"]
            add = {"text": rng.choice(pool).format(nm=_name(a)), "evidence_ids": []}
            return claims[:insert_at] + [add] + claims[insert_at:], f"causal/action claim on {_name(a)}: \"{add['text']}\""
        return None

    if etype == "omission_of_essential_driver":
        a = attrs[0]
        k = dc.get(a.eid)
        if k is None:
            return None
        rest = claims[:k] + claims[k + 1:]
        if _name(a).lower() in " ".join(c["text"] for c in rest).lower():
            return None
        return rest, f"strongest driver {_name(a)} ({a.eid}) removed: \"{claims[k]['text']}\""
    return None


def clean_insertion(claims, eo: EvidenceObject, rng: random.Random):
    """A CLEAN control written in the insertion frames: restates a real driver
    with its true direction, so insertion wording alone never signals an error."""
    attrs = sorted(eo.by_type(EvidenceType.MODEL_ATTRIBUTION), key=lambda a: -abs(a.contribution or 0))
    if len(attrs) < 2:
        return None
    a = rng.choice(attrs[1:])
    verb = "raises" if a.direction == "increases_risk" else "lowers"
    add = {"text": rng.choice(FRAMES["main"]["insert"]).format(nm=_name(a), verb=verb), "evidence_ids": [a.eid]}
    at = max(1, len(claims) - 2)
    return [dict(c) for c in claims[:at]] + [add] + [dict(c) for c in claims[at:]], \
        f"clean restatement of {_name(a)} in insertion wording (no error)"


PARA_STYLES = 3      # main-pack paraphrase styles; style 3 is practice-only


def paraphrase(claims, eo: EvidenceObject, rng: random.Random, style: Optional[int] = None):
    """A valid rewording: same facts and numbers, different phrasing/order."""
    attrs = sorted(eo.by_type(EvidenceType.MODEL_ATTRIBUTION), key=lambda a: -abs(a.contribution or 0))
    pred = eo.prediction
    if pred is None or not attrs:
        return None
    up = [_name(a) for a in attrs if a.direction == "increases_risk"]
    dn = [_name(a) for a in attrs if a.direction == "decreases_risk"]
    style = rng.randrange(PARA_STYLES) if style is None else style
    pdt = f"{pred.calibrated_pd:.1%}"
    disclaimer = None
    if style == 0:
        parts = [f"According to the model, the estimated probability of default is {pdt} for the documented outcome."]
        if up:
            parts.append(f"Relative to the model's reference, {', '.join(up)} {'push' if len(up) > 1 else 'pushes'} the risk score up.")
        if dn:
            parts.append(f"{', '.join(dn)} {'pull' if len(dn) > 1 else 'pulls'} the risk score down.")
        parts.append(f"The largest single contribution comes from {_name(attrs[0])}.")
        disclaimer = "These are associations inside the model, not established causes of default."
        parts.append(disclaimer)
    elif style == 1:
        parts = [f"The strongest driver of the score is {_name(attrs[0])}."]
        for a in attrs:
            parts.append(f"{_name(a)} {'raises' if a.direction == 'increases_risk' else 'lowers'} the model's risk score.")
        parts.append(f"The calibrated probability of default is {pdt}.")
        disclaimer = "The model describes associations; it does not show what causes default."
        parts.append(disclaimer)
    elif style == 2:
        parts = [f"The model gives a default probability of {pdt}."]
        if up:
            parts.append(f"Factors that increase the model's risk score: {', '.join(up)}.")
        if dn:
            parts.append(f"Factors that decrease it: {', '.join(dn)}.")
        parts.append(f"Of these, {_name(attrs[0])} contributes the most.")
    else:                                    # practice-only wording
        parts = [f"Estimated chance of default: {pdt}.",
                 f"Main driver: {_name(attrs[0])}, which {'raises' if attrs[0].direction == 'increases_risk' else 'lowers'} the model's risk score."]
        disclaimer = "This reflects the model, not a proven cause."
        parts.append(disclaimer)
    note = "valid paraphrase (no error)" + (f"; the sentence \"{disclaimer}\" is a correct disclaimer, "
                                             "not a causal claim" if disclaimer else "")
    return [{"text": p, "evidence_ids": []} for p in parts], note


# ---------------------------------------------------------------------------
def extra_borrowers(ds: str, exclude: set, n: int, rng: random.Random):
    """Evidence + template for n development borrowers NOT in the main pack, the
    v1 pilot or the superseded attempt (challenge bases and practice items)."""
    import joblib
    from run_m1 import PRIOR, SUPERSEDED, load_corpus, role_of
    from evidence_builder import GroupedShap, build_evidence_v2
    from generate import render_template_v2
    X, y, groups, _, orig = load_corpus(ds)
    sub = joblib.load(M1 / f"substrate_{ds}.joblib")          # the exact M1 model
    sg = GroupedShap(sub)
    pool = [i for i in sub.metrics["_test_index"]
            if role_of(f"{ds}-{orig[i]}") == "development"
            and f"{ds}-{orig[i]}" not in exclude | PRIOR | SUPERSEDED
            and str(groups.iloc[i]) != "other"]
    pool = sorted(pool)
    rng.shuffle(pool)
    out = []
    for i in pool[:n]:
        eo = build_evidence_v2(ds, f"{ds}-{orig[i]}", X.iloc[[i]], sub, sg, str(groups.iloc[i]))
        out.append((eo, render_template_v2(eo).claims))
    return out


def practice_items(rng: random.Random, exclude: set):
    """Exactly 5 practice items per dataset on borrowers outside the main pack:
    2 clean (template, paraphrase) + 3 single errors."""
    out = []
    for ds, kinds in (("taiwan", ["sign_reversal", "numeric_or_unit_error", "invented_factor"]),
                      ("home_credit", ["omission_of_essential_driver", "unsupported_causal_or_action",
                                       "unsupported_trend"])):
        cand = extra_borrowers(ds, exclude, 14, rng)
        cases = []
        if len(cand) >= 2:
            eo, cl = cand[0]                                   # clean template
            exclude.add(eo.borrower_id)
            cases.append((eo, "clean", cl, "N", "No error: a faithful template-style explanation."))
            eo2, cl2 = cand[1]                                 # clean paraphrase, another borrower
            exclude.add(eo2.borrower_id)
            pcl, pnote = paraphrase(cl2, eo2, rng, style=PARA_STYLES)   # practice-only wording
            cases.append((eo2, "clean_para", pcl, "N", f"No error: {pnote}."))
        for et in kinds:
            for eo, cl in cand[2:]:
                if eo.borrower_id in exclude:
                    continue
                m = mutate(cl, eo, et, rng, frames="practice")
                if m:
                    exclude.add(eo.borrower_id)
                    cases.append((eo, et, m[0], "Y", f"ANY = Y ({et}): {m[1]}."))
                    break
        if len(cases) != 5:
            raise SystemExit(f"practice: only {len(cases)} valid cases for {ds}")
        out += [{"dataset": ds, "eo": eo, "kind": k, "text": " ".join(c["text"] for c in cl),
                 "ANY": y, "answer": ans} for eo, k, cl, y, ans in cases]
    assert len(out) == 10
    return out


# ---------------------------------------------------------------------------
GUIDE = r"""# Milestone-1 Annotation Guidelines — Credit Explanation Verification (v2)

Changes since round 1 are marked **NEW**.

## 1. Your task

Each row shows **EVIDENCE** (what a credit-risk model produced for one borrower)
and an **EXPLANATION** written from it for a credit officer. Decide whether the
explanation says anything the evidence does not support. Assume the evidence is
correct and judge the explanation only against it. You are not judging the
model, the loan decision, or writing style.

**NEW — two datasets.** The `dataset` column says which one the row comes from:
Taiwan credit-card data or Home Credit loan applications. Field meanings are in
`FEATURE_GLOSSARY.md`, and every evidence line now also states its meaning,
unit and when it was recorded.

## 2. Reading the evidence

| Type | Meaning | The explanation MAY say |
|---|---|---|
| Observed fact `[f…]` | A recorded value, with meaning, unit and period | The value (sensibly rounded), its meaning |
| Model attribution `[a…]` | A feature RAISES or LOWERS the model's risk score; `value in [f…]` points to its recorded value | That it raises / lowers the risk score |
| Model prediction `[p…]` | The calibrated probability of default, its outcome and horizon | That estimate |
| Not available | Something that could not be computed | That no estimate is available |

Rules for reading:
1. **Attribution contributions are not percentages.** "+0.61 on the raw log-odds
   margin" means the feature pushes the score up. Writing "61% more risk" is a
   `numeric_or_unit_error`.
2. **NEW — encoded values.** A categorical value is shown with its documented
   label, e.g. `EDUCATION = 2 (university)`. A field name alone never tells you
   the borrower's category. Some codes are marked *undocumented*: an explanation
   that treats an undocumented code as if its meaning were certain has an error.
3. **NEW — missing is not zero.** `not recorded` means no value exists. Treating
   it as 0 is a `numeric_or_unit_error`.
4. **NEW — periods.** Taiwan repayment, bill and payment fields are monthly
   (April–September 2005). Home Credit `DAYS_` fields count days before the
   application (negative numbers; -12,000 days is about 32.9 years).

## 3. Labels (Y / N)

Mark **Y** if the explanation contains at least one instance.

- **invented_factor** — names a factor, fact or reason that is not in the evidence.
  A plain-language name for an evidence field is fine ("credit limit" for LIMIT_BAL).
- **sign_reversal** — says a feature raises risk when the evidence says it LOWERS it, or vice versa.
- **numeric_or_unit_error** — **NEW: check each number against the field it is attached to.**
  "AGE is 50,000" is an error even if 50,000 appears elsewhere in the evidence.
  Sensible rounding is fine (38.4% → 38%). Attribution numbers presented as percentages
  or probability changes are errors.
- **unsupported_trend** — **NEW wording.** Asserts a change over time that the
  evidence does not show. A comparison of two periods that are **both shown in the
  evidence** is allowed if it is stated correctly (e.g. "2 months late in August 2005 and
  paid duly in September 2005"). "The risk score is increased by X" describes the model,
  not a trend.
- **unsupported_causal_or_action** — claims a real-world cause, or that an action will
  change the outcome. "The model associates X with higher risk" is fine. A disclaimer
  such as "this describes the model, not a verified cause" is correct, not a causal claim.
- **omission_of_essential_driver** — leaves out the strongest driver (the attribution with
  the largest absolute contribution) or the probability estimate. Naming the driver only by
  a broad category ("repayment history") without identifying it counts as omission.
- **other_material_error** — anything else that would mislead a credit officer (wrong
  outcome or horizon, two features confused, an undocumented code stated as certain). Explain in `notes`.
- **ANY_MATERIAL_ERROR** — Y if any of the seven above is Y.

**NEW — UNRESOLVED.** Mark `UNRESOLVED = Y` when you cannot decide because the
evidence is insufficient, the rubric does not apply, or a feature is unclear. Choose
`unresolved_reason`. Still give your best guess in `ANY_MATERIAL_ERROR`. Do not use
UNRESOLVED for ordinary hard cases.

**NEW — reasonable_disagreement.** Y if a careful colleague could reasonably give a
different `ANY_MATERIAL_ERROR` label. This is about the item, not your confidence.

- **informative_for_credit_officer** — Y if a credit officer would learn the estimate and the main reason(s).
- **confidence** — High / Medium / Low.
- **error_quote** — the exact wrong words, if ANY = Y.
- **NEW — error_evidence_ids** — the evidence ids involved, e.g. `a2, f3`.
- **notes** — anything useful.

## 4. Borderline rules

1. Wording is lenient, facts are strict.
2. One wrong sentence is enough for Y.
3. Hedged claims still count ("income may have declined" is a trend claim).
4. Do not penalise correct disclaimers or statements that something is unavailable.
5. Some rows look alike or identical. Label each on its own.

## 5. Practical

- **NEW — practice gate.** Fill the PRACTICE sheet first and return it. Your main
  sheet is counted only if you get **at least 8 of the 10** ANY_MATERIAL_ERROR labels
  right **and fill ANY_MATERIAL_ERROR on every practice row** (a blank counts as not
  passing; if you mark UNRESOLVED, still give your best guess). You will get the
  practice answers after you return it.
- About 1–2 minutes per row after the first 20. Take a break every 45 minutes.
- Work independently. Do not use ChatGPT or any AI tool.
- A broken row (empty or unreadable): ANY = Y, other_material_error = Y, note "broken".
"""

README = """SAARTHI — Milestone-1 Explanation Annotation ({rater})
=====================================================

Your annotator code: {rater}
Items in your main sheet: {n}   (practice sheet: {p})

Files
  README.txt                       this file
  ANNOTATION_GUIDELINES.md         READ FULLY first (changes since round 1 are marked NEW)
  FEATURE_GLOSSARY.md              what each field means, both datasets
  M1_PRACTICE_{rater}.xlsx         do this FIRST and return it
  M1_annotation_sheet_{rater}.xlsx the main work: fill the yellow columns

Steps
  1. Read ANNOTATION_GUIDELINES.md and FEATURE_GLOSSARY.md (about 30 minutes).
  2. Fill M1_PRACTICE_{rater}.xlsx and send it back. You will receive the answers.
     Your main sheet counts only if you get at least 8 of 10 practice rows right, with
     ANY_MATERIAL_ERROR filled on every practice row.
  3. Fill the Annotator tab (your code only) and label every row of the main sheet.
  4. Return the main sheet unchanged in name.

Rules
  - Work independently; do not discuss items until everyone has returned.
  - Do not use ChatGPT or any AI tool.
  - The material is for this study only. Do not share it.
"""


def write_sheet(path: Path, rows: list, rater: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Annotate"
    head = ["item_id", "dataset", "evidence", "explanation"] + LABEL_COLS
    ws.append(head)
    for r in rows:
        ws.append([r["item_id"], r["dataset"], r["evidence"], r["explanation"]] + [""] * len(LABEL_COLS))
    hdr, lab = PatternFill("solid", fgColor="1F2937"), PatternFill("solid", fgColor="FEF3C7")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), hdr
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, w in {"A": 11, "B": 14, "C": 80, "D": 70}.items():
        ws.column_dimensions[col].width = w
    for j in range(5, 5 + len(LABEL_COLS)):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 15
    for name in ("error_quote", "notes", "unresolved_reason"):
        ws.column_dimensions[ws.cell(row=1, column=5 + LABEL_COLS.index(name)).column_letter].width = 32
    for row in ws.iter_rows(min_row=2):
        for c in row[:4]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[4:]:
            c.fill, c.alignment = lab, Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    n = len(rows) + 1
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    conf = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True)
    why = DataValidation(type="list", allow_blank=True, formula1=
                         '"insufficient evidence,rubric not applicable,unclear feature,other"')
    for dv in (yn, conf, why):
        ws.add_data_validation(dv)
    col = lambda nm: ws.cell(row=1, column=5 + LABEL_COLS.index(nm)).column_letter
    for nm in YN_COLS:
        yn.add(f"{col(nm)}2:{col(nm)}{n}")
    conf.add(f"{col('confidence')}2:{col('confidence')}{n}")
    why.add(f"{col('unresolved_reason')}2:{col('unresolved_reason')}{n}")
    meta = wb.create_sheet("Annotator")
    for r in [["annotator_code", rater], ["date started", ""], ["date finished", ""],
              ["total time spent (hours)", ""], ["used any AI tool? (must be N)", ""]]:
        meta.append(r)
    meta.column_dimensions["A"].width, meta.column_dimensions["B"].width = 40, 24
    wb.save(path)


# ---------------------------------------------------------------------------
def main(force: bool = False) -> None:
    import hashlib
    if (SUB / "annotation_key_m1.jsonl").exists() and not force:
        raise SystemExit("annotation_key_m1.jsonl exists: the pack was already built (and maybe sent). "
                         "Rebuilding would change the item-id mapping. Use --force only if nothing was sent.")
    rng = random.Random(SEED)
    raw_bytes = (M1 / "records.jsonl").read_bytes()
    # records.jsonl is append-only: the CURRENT row per item is the last one
    latest = {}
    for l in raw_bytes.decode().splitlines():
        if l.strip():
            r = json.loads(l)
            latest[(r["borrower_id"], r["arm"], r["family"])] = r
    recs = list(latest.values())
    natural = sorted((r for r in recs if not r.get("error") and (r.get("text") or "").strip()),
                     key=lambda r: (r["borrower_id"], r["arm"], r["family"]))   # stable order
    # half of the challenge items sit on borrowers that also have natural items
    # (evidence repetition carries no signal); that borrower's clean template is
    # withheld from the pack so no rater can compare a mutation with its twin
    tmpl = {r["borrower_id"]: r for r in natural if r["arm"] == "template"}
    host = {}
    for ds in sorted({r["dataset"] for r in natural}):
        cands = sorted(b for b, r in tmpl.items() if r["dataset"] == ds)
        rng.shuffle(cands)
        host[ds] = cands[:12]
    withheld = {b for v in host.values() for b in v}
    natural = [r for r in natural if not (r["arm"] == "template" and r["borrower_id"] in withheld)]

    items, key = [], []
    for r in natural:                                      # natural outputs
        eo = eo_from(r["evidence"])
        items.append({"eo": eo, "text": blind(r["text"], eo), "dataset": r["dataset"]})
        key.append({"source": "natural", "borrower_id": r["borrower_id"], "dataset": r["dataset"],
                    "group": r["group"], "pd_band": r["pd_band"], "arm": r["arm"], "family": r["family"],
                    "model_id": r["model_id"], "route_index": r["route_index"],
                    "previously_inspected": r.get("previously_inspected", False),
                    "evidence_fingerprint": r["evidence_fingerprint"], "injected_error": None})

    # challenge items: half on natural borrowers (template withheld), half on
    # HELD-OUT development borrowers; each base borrower carries one challenge item
    used = {r["borrower_id"] for r in recs}
    types = ["invented_factor", "sign_reversal", "numeric_or_unit_error", "unsupported_trend",
             "unsupported_causal_or_action", "omission_of_essential_driver"]
    for ds in sorted({r["dataset"] for r in natural}):
        hosted = [(eo_from(tmpl[b]["evidence"]), tmpl[b]["claims"]) for b in host.get(ds, [])]
        held = extra_borrowers(ds, used, 80, rng)
        bases = [x for pair in zip(hosted, held[:len(hosted)]) for x in pair] + held[len(hosted):]
        host_ids = set(host.get(ds, []))
        used -= host_ids                     # hosted borrowers may carry one challenge item
        # per dataset: 3 x 6 single-error mutations + 3 clean paraphrases + 3 clean
        # insertions (insertion wording alone never signals an error)
        plan = [t for t in types for _ in range(3)] + ["paraphrase"] * 3 + ["clean_insertion"] * 3
        for kind in plan:
            got = None
            for eo, cl in bases:
                if eo.borrower_id in used:
                    continue
                gen = (paraphrase(cl, eo, rng) if kind == "paraphrase" else
                       clean_insertion(cl, eo, rng) if kind == "clean_insertion" else
                       mutate(cl, eo, kind, rng))
                if gen:
                    got = (eo, gen)
                    break
            if got is None:
                raise SystemExit(f"{ds}: no base borrower left for {kind}")
            eo, (out_claims, note) = got
            used.add(eo.borrower_id)
            text = " ".join(c["text"] for c in out_claims)
            items.append({"eo": eo, "text": blind(text, eo), "dataset": ds})
            key.append({"source": "challenge", "borrower_id": eo.borrower_id, "dataset": ds,
                        "base": "natural_borrower" if eo.borrower_id in host_ids else "held_out",
                        "group": eo.group_attributes.get("audit_group"),
                        "arm": "template+" + kind if kind in ("paraphrase", "clean_insertion") else "template+mutation",
                        "family": "template", "challenge_kind": kind,
                        "injected_error": None if kind in ("paraphrase", "clean_insertion") else kind,
                        "mutation_note": note, "evidence_fingerprint": eo.fingerprint()})

    # balanced pair assignment within dataset x source strata
    order = list(range(len(items)))
    rng.shuffle(order)
    strata = defaultdict(list)
    for i in order:   # challenge items stratified by kind: each rater sees 2 of 3 per type
        strata[(key[i]["dataset"], key[i]["source"], key[i].get("challenge_kind") or "")].append(i)
    assign, offset = {}, 0
    for s in sorted(strata):
        for n, i in enumerate(strata[s]):
            assign[i] = PAIRS[(n + offset) % 3]
        offset += len(strata[s])
    per_rater = {r: [i for i in order if r in assign[i]] for r in RATERS}

    # hidden duplicates; ids for originals AND duplicates from ONE shuffled space
    slots = [("item", i, None) for i in order]
    for r in RATERS:
        for i in rng.sample(per_rater[r], max(1, round(DUP_RATE * len(per_rater[r])))):
            slots.append(("dup", i, r))
    ids = [f"M1-{n:04d}" for n in range(1, len(slots) + 1)]
    rng.shuffle(ids)
    iid = {}
    for (kind, i, r), sid in zip(slots, ids):
        if kind == "item":
            iid[i] = sid
    for k, it in zip(key, items):
        k["text_sha256"] = hashlib.sha256(it["text"].encode()).hexdigest()[:16]
    keyrows, dup_rows = [], {r: [] for r in RATERS}
    for (kind, i, r), sid in zip(slots, ids):
        if kind == "item":
            keyrows.append({"item_id": sid, "raters": list(assign[i]), **key[i]})
        else:
            dup_rows[r].append((sid, i))
            keyrows.append({"item_id": sid, "raters": [r], "duplicate_of": iid[i], **key[i]})

    # practice: borrowers outside the pack, the challenge bases and v1/superseded
    prac = practice_items(rng, set(used) | {r["borrower_id"] for r in recs})
    rng.shuffle(prac)
    pkey = {}
    for n, p in enumerate(prac, 1):
        pkey[f"P-{n:02d}"] = {"ANY": p["ANY"], "kind": p["kind"], "answer": p["answer"],
                              "borrower_id": p["eo"].borrower_id}

    for f in SUB.glob("annotation_pack_m1_*"):
        shutil.rmtree(f) if f.is_dir() else f.unlink()
    glossary = FS.glossary_markdown()
    load = {}
    for r in RATERS:
        pack = SUB / f"annotation_pack_m1_{r}"
        pack.mkdir(parents=True)
        rows = [(iid[i], i) for i in per_rater[r]] + dup_rows[r]
        rng.shuffle(rows)
        sheet = [{"item_id": rid, "dataset": DATASET_LABEL[items[i]["dataset"]],
                  "evidence": render_evidence_v2(items[i]["eo"]), "explanation": items[i]["text"]}
                 for rid, i in rows]
        all_bids = {k["borrower_id"] for k in key}
        for row in sheet:
            assert_blind(row["evidence"] + " " + row["explanation"], all_bids)
        write_sheet(pack / f"M1_annotation_sheet_{r}.xlsx", sheet, r)
        psheet = [{"item_id": pid, "dataset": DATASET_LABEL[p["dataset"]],
                   "evidence": render_evidence_v2(p["eo"]), "explanation": blind(p["text"], p["eo"])}
                  for pid, p in zip(sorted(pkey), prac)]
        write_sheet(pack / f"M1_PRACTICE_{r}.xlsx", psheet, r)
        (pack / "ANNOTATION_GUIDELINES.md").write_text(GUIDE)
        (pack / "FEATURE_GLOSSARY.md").write_text(glossary)
        (pack / "README.txt").write_text(README.format(rater=r, n=len(sheet), p=len(prac)))
        zp = SUB / f"annotation_pack_m1_{r}.zip"
        with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(pack.iterdir()):
                z.write(f, f"annotation_pack_m1_{r}/{f.name}")
        load[r] = len(sheet)

    meta = {"_meta": True, "records_sha256": hashlib.sha256(raw_bytes).hexdigest()[:16],
            "withheld_templates": sorted(withheld), "seed": SEED}
    (SUB / "annotation_key_m1.jsonl").write_text(
        "\n".join(json.dumps(k, default=str) for k in [meta] + keyrows) + "\n")
    (SUB / "practice_key_m1.json").write_text(json.dumps(
        {"pass_mark": PASS_MARK, "n": len(prac), "items": pkey}, indent=2))
    c = Counter((k["source"], k["dataset"]) for k in keyrows if "duplicate_of" not in k)
    inj = Counter(k["injected_error"] for k in keyrows if k["source"] == "challenge" and "duplicate_of" not in k)
    print(f"unique items={len(order)} {dict(c)}\nchallenge types={dict(inj)}")
    print(f"per-rater main sheet sizes={load} (each item rated by 2 of 3)")
    print(f"practice items={len(prac)} kinds={[p['kind'] for p in prac]}")
    print("send: " + ", ".join(f"annotation_pack_m1_{r}.zip" for r in RATERS))
    print("KEEP LOCAL: annotation_key_m1.jsonl, practice_key_m1.json")


if __name__ == "__main__":
    main(force="--force" in sys.argv)
