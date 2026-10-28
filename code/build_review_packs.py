"""
Two human-review packs for label independence (Improvement_Plan WS3).

1. Independent_rating_pack: a THIRD, independent rating of locked-test items, by a person who is not an author and
   who did not rate these items. Blind to earlier labels, adjudications, verifier outputs and arm.
     - every item whose two ratings disagreed on material error (adjudicated before by the first author)
     - every item whose label notes record a revision after feedback
     - a random audit sample of agreed items (50 narratives, 20 templates)
   Same sheet and guidelines as the locked test.
2. Context_check_pack: a human truth check of the controlled-context focal sentences (packs v1 and v2): every
   NEGATION and REFERENCE focal sentence (correct and single-error versions) and 40 random other focal sentences.
   The checker sees the evidence and the sentence and answers whether the sentence is correct; the intended label
   is hidden.

    python build_review_packs.py
Outputs: Submission/packs/Independent_rating_pack.zip, Submission/packs/Context_check_pack.zip (send)
         Submission/annotation_key_independent_v1.jsonl, Submission/mechanism/annotation_key_context_check.jsonl (LOCAL)
"""
from __future__ import annotations

import csv
import json
import random
import re
import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_locked_pack as BL  # noqa: E402
import build_pack_m1 as PM  # noqa: E402
import feature_semantics as FS  # noqa: E402
from build_pack_m1 import SUB, assert_blind, blind, eo_from  # noqa: E402
from evidence_builder import render_evidence_v2  # noqa: E402
from score_m1 import read_sheet, yn  # noqa: E402

SEED = 20261014
PACKS = SUB / "packs"
REV = re.compile(r"supervisor|corrected|updated|revised|feedback", re.I)

IND_README = """Independent rating ({rater})
=====================

You are an INDEPENDENT rater. You must not be an author of the study and you must not have labelled these
explanations before. Earlier labels and any software verdicts are hidden from you on purpose.

Files
  README.txt                       this file
  ANNOTATION_GUIDELINES.md         the same rules the first raters used (read the addendum at the end)
  FEATURE_GLOSSARY.md              what each field means
  INDEPENDENT_rating_sheet.xlsx    fill the yellow columns for every row

Rules
  - Work alone; do not discuss items with the authors or the earlier raters until you return the sheet.
  - Do not use ChatGPT or any AI tool.
  - The material is for this study only.
"""

CHECK_README = """Context sentence check
======================

Each row shows one borrower's evidence and ONE sentence about it. Decide whether the sentence is CORRECT with
respect to the evidence. Read negations carefully ('is not used to lower', 'does not raise').

Columns to fill (yellow):
  SENTENCE_CORRECT   Y if everything the sentence states is supported by the evidence; N otherwise
  ERROR_TYPE         if N: value | direction (incl. negation) | strongest driver | scale (log-odds vs %) |
                     fact given a direction | change over time | other
  UNCLEAR            Y if the sentence is ambiguous and you cannot decide
  NOTES              optional

Contributions are on the raw log-odds scale; a field listed only as a recorded fact (no driver entry) has no
direction of its own. Do not use any AI tool. Do not discuss items with the authors before returning the sheet.
"""


def independent_pack(rng):
    key = [json.loads(l) for l in open(SUB / "annotation_key_locked_v1.jsonl")][1:]
    D = SUB / "Annotated" / "IDBI_Hack_Annotation"
    lab = {q: read_sheet(sorted(D.rglob(f"LOCKED_annotation_sheet_{q}*.xlsx"))[-1])["rows"] for q in BL.RATERS}
    adj = {r["item_id"] for r in csv.DictReader(open(SUB / "locked" / "adjudication_locked.csv"))}
    dup_of = {k["item_id"]: k["duplicate_of"] for k in key if k.get("duplicate_of")}
    revised = set()
    for q, rows in lab.items():
        for i, r in rows.items():
            if REV.search(str(r.get("notes") or "")):
                revised.add(dup_of.get(i, i))
    uniq = {k["item_id"]: k for k in key if not k.get("duplicate_of")}
    chosen = {}
    for i in adj:
        chosen[i] = "disagreement"
    for i in revised:
        chosen.setdefault(i, "revised_after_feedback")
    agreed = [i for i in uniq if i not in chosen]
    nar = sorted(i for i in agreed if uniq[i]["arm"] == "grounded")
    tmp = sorted(i for i in agreed if uniq[i]["arm"] == "template")
    for i in rng.sample(nar, 50):
        chosen[i] = "agreed_audit"
    for i in rng.sample(tmp, 20):
        chosen[i] = "agreed_audit"
    recs = {json.loads(l)["item_id"]: json.loads(l) for l in open(SUB / "locked" / "fresh_natural_v1.jsonl")}
    items = sorted(chosen)
    rng.shuffle(items)
    sheet, keyrows = [], []
    for n, i in enumerate(items, 1):
        k = uniq[i]
        r = recs[k["source_item_id"]]
        eo = eo_from(r["evidence"])
        sid = f"IR-{n:03d}"
        sheet.append({"item_id": sid, "dataset": PM.DATASET_LABEL[r["dataset"]], "evidence": render_evidence_v2(eo),
                      "explanation": blind(r["text"], eo)})
        keyrows.append({"item_id": sid, "locked_item_id": i, "source_item_id": k["source_item_id"],
                        "reason": chosen[i], "arm": k["arm"], "borrower_id": k["borrower_id"]})
    all_b = {k["borrower_id"] for k in keyrows}
    for row in sheet:
        assert_blind(row["evidence"] + " " + row["explanation"], all_b)
    stage = PACKS / "Independent_rating_pack"
    shutil.rmtree(stage, ignore_errors=True)
    stage.mkdir(parents=True)
    old = list(PM.LABEL_COLS)
    PM.LABEL_COLS[:] = old[:-3] + BL.EXTRA + old[-3:]
    try:
        PM.write_sheet(stage / "INDEPENDENT_rating_sheet.xlsx", sheet, "IR")
    finally:
        PM.LABEL_COLS[:] = old
    (stage / "ANNOTATION_GUIDELINES.md").write_text(PM.GUIDE + BL.ADDENDUM)
    (stage / "FEATURE_GLOSSARY.md").write_text(FS.glossary_markdown())
    (stage / "README.txt").write_text(IND_README.format(rater="independent rater"))
    z = PACKS / "Independent_rating_pack.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(stage.iterdir()):
            zf.write(f, f"Independent_rating_pack/{f.name}")
    (SUB / "annotation_key_independent_v1.jsonl").write_text("\n".join(json.dumps(x) for x in keyrows) + "\n")
    from collections import Counter
    print(z, len(sheet), dict(Counter(x["reason"] for x in keyrows)))


def context_pack(rng):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    rows = []
    for tag in ("v1", "v2"):
        for l in open(SUB / "mechanism" / f"context_pack_{tag}.jsonl"):
            r = json.loads(l)
            if r["context"] == "alone":
                rows.append(dict(r, pack=tag))
    focus = [r for r in rows if r["relation"] in ("NEGATION", "REFERENCE")]
    rest = [r for r in rows if r["relation"] not in ("NEGATION", "REFERENCE")]
    pick = focus + rng.sample(rest, 40)
    rng.shuffle(pick)
    wb = Workbook()
    ws = wb.active
    ws.title = "Check"
    ws.append(["item_id", "dataset", "evidence", "sentence", "SENTENCE_CORRECT", "ERROR_TYPE", "UNCLEAR", "NOTES"])
    keyrows = []
    for n, r in enumerate(pick, 1):
        eo = eo_from(r["evidence"])
        sid = f"CC-{n:03d}"
        ws.append([sid, PM.DATASET_LABEL[r["dataset"]], render_evidence_v2(eo), blind(r["focal_text"], eo), "", "", "", ""])
        keyrows.append({"item_id": sid, "pack": r["pack"], "context_item_id": r["item_id"], "base_id": r["base_id"],
                        "relation": r["relation"], "intended_error": r["truth"], "borrower_id": r["borrower_id"]})
    hdr, fill = PatternFill("solid", fgColor="1F2937"), PatternFill("solid", fgColor="FEF3C7")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), hdr
    for col, w in zip("ABCDEFGH", (9, 14, 80, 60, 12, 18, 9, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[4:]:
            c.fill = fill
    ws.freeze_panes = "E2"
    meta = wb.create_sheet("Checker")
    for r in [["checker_code", ""], ["used any AI tool? (must be N)", ""]]:
        meta.append(r)
    bids = {k["borrower_id"] for k in keyrows}
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert_blind(" ".join(str(x) for x in row if x), bids)
    stage = PACKS / "Context_check_pack"
    shutil.rmtree(stage, ignore_errors=True)
    stage.mkdir(parents=True)
    wb.save(stage / "CONTEXT_check_sheet.xlsx")
    (stage / "README.txt").write_text(CHECK_README)
    (stage / "FEATURE_GLOSSARY.md").write_text(FS.glossary_markdown())
    z = PACKS / "Context_check_pack.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(stage.iterdir()):
            zf.write(f, f"Context_check_pack/{f.name}")
    (SUB / "mechanism" / "annotation_key_context_check.jsonl").write_text("\n".join(json.dumps(x) for x in keyrows) + "\n")
    from collections import Counter
    print(z, len(keyrows), dict(Counter((x["relation"], x["intended_error"]) for x in keyrows)))


if __name__ == "__main__":
    rng = random.Random(SEED)
    independent_pack(rng)
    context_pack(rng)
