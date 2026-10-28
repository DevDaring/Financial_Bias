"""
Practitioner task study, pilot (Improvement_Plan WS5). Credit/risk practitioners see what a lender would actually show
for a borrower under three policies and answer factual questions; correctness is scored against the evidence.

Conditions (what is shown for a case):
  T  the audited template
  G  the narrative if the strongest generic verifier (B1F, GPT-OSS judge) released it, otherwise the template
  R  the narrative if ReCheck released it, otherwise the template
Cases: locked-v1 borrowers (labels known; development data), stratified so that a third of the cases have a WRONG
narrative that at least one verifier released. Each participant sees 15 cases, 5 per condition, in a Latin-square
rotation (no borrower twice for one participant). The output's source is not revealed.

Questions per case (scored against the evidence):
  Q1 the model's probability of default (value)
  Q2 the strongest driver and whether it raises or lowers risk
  Q3 does the text contain a statement that is not supported by the evidence? (Y/N; quote it)
  Q4 would you ask for clarification before using this text with a customer? (Y/N)
  Q5 time taken (the participant records start/end minutes in the sheet)
Primary outcomes: Q1-Q2 accuracy, Q3 detection of a misleading statement, time; Q4 and usefulness are secondary.

    python build_practitioner_pack.py --participants 12
Outputs: Submission/packs/Practitioner_pilot_pack.zip (send one sheet per participant)
         Submission/practitioner_key_v1.jsonl (LOCAL ONLY)
"""
from __future__ import annotations

import argparse
import json
import random
import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_pack_m1 as PM  # noqa: E402
from build_pack_m1 import SUB, assert_blind, blind, eo_from  # noqa: E402
from dev_sets import locked_items  # noqa: E402
from evidence_builder import render_evidence_v2  # noqa: E402

SEED = 20261018
README = """Credit-explanation task (pilot)
==============================

Thank you for taking part. Each case shows a borrower's model evidence and a short text that a bank might show
about the model's assessment. For each case answer the questions in the yellow columns. Record the clock time when
you start and finish each case. There are no trick questions about the evidence itself; read the text carefully.

Columns
  START_TIME / END_TIME     clock time (hh:mm) when you start and finish the case
  Q1_PD_VALUE               the probability of default the model estimates (as a %), from the TEXT if it states it,
                            otherwise from the evidence
  Q2_STRONGEST_DRIVER       the field that drives the model's score most strongly
  Q2_DIRECTION              does it RAISE or LOWER the risk score?
  Q3_UNSUPPORTED (Y/N)      does the TEXT state anything the evidence does not support?
  Q3_QUOTE                  if Y, the words
  Q4_CLARIFY (Y/N)          would you ask for clarification before using this text with a customer?
  USEFUL (1-5)              how useful the text is for explaining the decision

Rules: work alone; do not use any AI tool; about 2-3 minutes per case.
Consent: by returning this sheet you agree that your answers (without your name) may be used for research.
"""


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--participants", type=int, default=12)
    ap.add_argument("--per-condition", type=int, default=5)
    a = ap.parse_args()
    rng = random.Random(SEED)
    items = {x["item_id"]: x for x in locked_items()}
    runs = SUB / "locked" / "fresh_narratives_v1_runs"
    g = {json.loads(l)["item_id"]: json.loads(l).get("accepted") for l in open(runs / "facts_judge__gptoss120b.jsonl")}
    rp = SUB / "locked" / "fresh_narratives_v1_runs_v16" / "exec_P__gptoss120b.jsonl"
    r = {json.loads(l)["item_id"]: json.loads(l).get("accepted") for l in open(rp)}
    tmpl = {}
    for l in open(SUB / "locked" / "fresh_natural_v1.jsonl"):
        x = json.loads(l)
        if x["arm"] == "template":
            tmpl[x["borrower_id"]] = x
    wrong_rel = [i for i, x in items.items() if x["truth"] and (g.get(i) is True or r.get(i) is True)]
    correct = [i for i, x in items.items() if not x["truth"]]
    n_cases = a.per_condition * 3
    pool = rng.sample(sorted(wrong_rel), min(len(wrong_rel), 2 * n_cases // 3 + 4)) + \
        rng.sample(sorted(correct), 2 * n_cases)
    pool = pool[:3 * n_cases]
    rng.shuffle(pool)
    stage = SUB / "packs" / "Practitioner_pilot_pack"
    shutil.rmtree(stage, ignore_errors=True)
    stage.mkdir(parents=True)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    key = []
    cols = ["case_id", "dataset", "evidence", "text", "START_TIME", "END_TIME", "Q1_PD_VALUE", "Q2_STRONGEST_DRIVER",
            "Q2_DIRECTION", "Q3_UNSUPPORTED (Y/N)", "Q3_QUOTE", "Q4_CLARIFY (Y/N)", "USEFUL (1-5)"]
    for p in range(a.participants):
        cases = pool[(p * n_cases) % len(pool):][:n_cases]
        if len(cases) < n_cases:
            cases += pool[:n_cases - len(cases)]
        wb = Workbook()
        ws = wb.active
        ws.title = "Cases"
        ws.append(cols)
        for k, i in enumerate(cases):
            cond = "TGR"[(k + p) % 3]
            x = items[i]
            t = tmpl[x["borrower_id"]]
            shown = {"T": t, "G": x["rec"] if g.get(i) is True else t, "R": x["rec"] if r.get(i) is True else t}[cond]
            eo = eo_from(x["rec"]["evidence"])
            cid = f"P{p + 1:02d}-C{k + 1:02d}"
            ws.append([cid, PM.DATASET_LABEL[x["dataset"]], render_evidence_v2(eo), blind(shown["text"], eo)] + [""] * 9)
            key.append({"case_id": cid, "participant": f"P{p + 1:02d}", "condition": cond, "item_id": i,
                        "borrower_id": x["borrower_id"], "shown": "narrative" if shown is x["rec"] else "template",
                        "narrative_wrong": x["truth"], "shown_wrong": bool(x["truth"]) if shown is x["rec"] else False})
        hdr, fill = PatternFill("solid", fgColor="1F2937"), PatternFill("solid", fgColor="FEF3C7")
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), hdr
        for col, w in zip("ABCDEFGHIJKLM", (11, 14, 70, 60, 10, 10, 12, 22, 12, 12, 30, 12, 10)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for c in row[4:]:
                c.fill = fill
        meta = wb.create_sheet("Participant")
        for rr in [["role (e.g. credit officer, risk analyst)", ""], ["years of experience", ""],
                   ["used any AI tool? (must be N)", ""]]:
            meta.append(rr)
        for row in ws.iter_rows(min_row=2, values_only=True):
            assert_blind(" ".join(str(v) for v in row if v), {k["borrower_id"] for k in key})
        wb.save(stage / f"PRACTITIONER_task_P{p + 1:02d}.xlsx")
    (stage / "README.txt").write_text(README)
    z = SUB / "packs" / "Practitioner_pilot_pack.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(stage.iterdir()):
            zf.write(f, f"Practitioner_pilot_pack/{f.name}")
    (SUB / "practitioner_key_v1.jsonl").write_text("\n".join(json.dumps(k) for k in key) + "\n")
    from collections import Counter
    print(z, a.participants, "participants x", n_cases, "cases;", Counter((k["condition"], k["shown"], k["shown_wrong"]) for k in key))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
