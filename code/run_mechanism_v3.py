"""
Held-out mechanism test v3: sentences written by an independent annotator (AB, 7 Oct 2026) from
Heldout_writing_pack, scored ONCE with the checker frozen at executable/1.4.

Items are single claims (original, paraphrase, corruption, provocation of 60 base claims on 30 fresh
borrowers). Labels are the writer's intent (corruption = error, otherwise correct); a second rater's
blind validation is still pending and is reported as a limitation.

Methods (claim level: omissions are not errors for a single sentence):
  A0  code layer only                      (no LLM)
  B1  repaired generic judge (judge/2.0)   generic_judge2
  B2  generic judge + same tools           generic_tool_judge
  P   executable verification              exec_P, claim-level mode
on --judge qwen3_next and/or grok45. DRY RUN unless --live --confirm-spend.

    python run_mechanism_v3.py --build                 # pack from the returned workbook + local key
    python run_mechanism_v3.py --judges qwen3_next grok45 --live --confirm-spend --workers 12
Output: Submission/mechanism/mechanism_pack_v3_annotator.jsonl, Submission/mechanism/v3_runs/*.jsonl,
        Submission/mechanism/eval_v3_once.json (refuses to overwrite)
"""
from __future__ import annotations

import argparse
import json
import sys
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import comparators as K  # noqa: E402
import executable_verifier as EV  # noqa: E402
import fakes  # noqa: E402
import witness as W  # noqa: E402
from build_pack_m1 import SUB, eo_from  # noqa: E402
from generate import Candidate  # noqa: E402

MECH = SUB / "mechanism"
PACK = MECH / "mechanism_pack_v3_annotator.jsonl"
RUNS = MECH / "v3_runs"
WB = SUB / "packs" / "Annotated" / "Heldout_writing_workbook_AB.xlsx"
KEY = MECH / "annotation_key_mechanism_v3.jsonl"


def build() -> int:
    from openpyxl import load_workbook
    if PACK.exists():
        print(f"{PACK} exists (frozen)")
        return 1
    key = {json.loads(l)["borrower_code"]: json.loads(l) for l in open(KEY)}
    ws = load_workbook(WB, data_only=True)["Write"]
    hdr = [c.value for c in ws[1]]
    meta = dict(zip(hdr[0::2], hdr[1::2]))
    rows = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r[0] or not r[7]:
            continue
        item, bcode, _ds, rel, targets, ver, _what, text, err, obj = r[:10]
        k = key[bcode]
        rows.append({"item_id": item, "base_id": item.rsplit("-", 1)[0], "borrower_code": bcode,
                     "borrower_id": k["borrower_id"], "dataset": k["dataset"],
                     "group": (k["evidence"].get("group_attributes") or {}),
                     "relation": rel, "version": ver, "text": str(text).strip(),
                     "label": "error" if ver == "corruption" else "correct",
                     "intended_error": err, "invited_objection": obj, "target_ids": targets,
                     "writer": meta.get("annotator_code"), "writer_hours": meta.get("hours"),
                     "writer_ai_tool": meta.get("used_AI_tool (must be N)"),
                     "labels_basis": "writer intent; blind second-rater validation pending",
                     "evidence": k["evidence"]})
    with open(PACK, "w") as f:
        for x in rows:
            f.write(json.dumps(x, default=str) + "\n")
    print(f"{PACK}: {len(rows)} items, {len({x['base_id'] for x in rows})} base claims",
          dict(Counter(x["relation"] for x in rows if x["version"] == "original")))
    return 0


def claim_exec(text, eo):
    tx = W.execute_text(text, eo)
    st = [c.status for c in tx.claims]
    if W.CONTRA in st:
        return False, "executed_contradiction", tx
    if st and all(s in (W.SUPPORT, "nonmaterial") for s in st) and W.SUPPORT in st:
        return True, "all_claims_executed_supported", tx
    return None, "residual", tx


def run_one(method, item, judge):
    eo = eo_from(item["evidence"])
    cand = Candidate(borrower_id=item["borrower_id"], arm="heldout_v3", text=item["text"], claims=[], family="")
    if method == "exec_A0":
        acc, why, _ = claim_exec(item["text"], eo)
        return {"accepted": acc, "reason": why, "calls": []}
    if method == "exec_P":
        v = EV.verify_executable(cand, eo, judge, mode="P", reassess_family=judge, experiment="mechanism-v3",
                                 claim_level=True)
        return {"accepted": v.accepted, "reason": v.reason, "calls": v.calls,
                "witnesses": [{k: w.get(k) for k in ("rule", "outcome", "quote")} for w in v.witnesses]}
    fn = K.METHODS[method]
    v = fn(cand, eo, family=judge, experiment="mechanism-v3")
    return {"accepted": v.accepted, "reason": v.error or "", "calls": v.calls,
            "violations": (v.detail or {}).get("violations")}


def score(rows):
    def frac(sel, cond):
        return f"{sum(1 for r in sel if cond(r))}/{len(sel)}"
    err = [r for r in rows if r["label"] == "error"]
    ok = [r for r in rows if r["label"] == "correct"]
    prov = [r for r in ok if r["version"] == "provocation"]
    return {"n": len(rows), "errors_rejected": frac(err, lambda r: r["accepted"] is False),
            "errors_accepted": frac(err, lambda r: r["accepted"] is True),
            "correct_accepted": frac(ok, lambda r: r["accepted"] is True),
            "correct_rejected": frac(ok, lambda r: r["accepted"] is False),
            "deferred": sum(1 for r in rows if r["accepted"] is None),
            "provocations_rejected": frac(prov, lambda r: r["accepted"] is False),
            "calls": sum(len(r.get("calls") or []) for r in rows)}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--build", action="store_true")
    ap.add_argument("--judges", nargs="*", default=["qwen3_next"])
    ap.add_argument("--methods", nargs="*", default=["exec_A0", "generic_judge2", "generic_tool_judge", "exec_P"])
    ap.add_argument("--live", action="store_true")
    ap.add_argument("--confirm-spend", action="store_true")
    ap.add_argument("--workers", type=int, default=1)
    ap.add_argument("--call-timeout", type=float, default=None)
    a = ap.parse_args()
    if a.call_timeout:
        import gateway as _G
        _G.GATEWAY.timeout = a.call_timeout
    if a.build:
        return build()
    items = [json.loads(l) for l in open(PACK)]
    if a.live and not a.confirm_spend:
        print("live run needs --confirm-spend")
        return 2
    RUNS.mkdir(parents=True, exist_ok=True)
    tag = "" if a.live else "__dry"
    ctx = fakes.patched(fakes.FakeGateway(lambda f, m: '{"accepted": true, "violations": []}'
                                          if "code_status" not in str(m[-1]["content"])
                                          else json.dumps({"claims": [], "omissions": []}))) if not a.live else None
    if ctx:
        ctx.__enter__()
    summary = {}
    for judge in a.judges:
        for m in a.methods:
            jn = "none" if m == "exec_A0" else judge
            if m == "exec_A0" and judge != a.judges[0]:
                continue
            path = RUNS / f"{m}__{jn}{tag}.jsonl"
            done = {json.loads(l)["item_id"]: json.loads(l) for l in open(path)} if path.exists() else {}
            rows = list(done.values())
            lock = threading.Lock()

            def work(it, m=m, judge=judge, path=path):
                t0 = time.time()
                from run_comparators import with_deadline
                out = with_deadline(lambda: run_one(m, it, judge), 600.0)
                out.setdefault("reason", out.get("error", ""))
                row = {"item_id": it["item_id"], "method": m, "judge": jn, "label": it["label"],
                       "version": it["version"], "relation": it["relation"], "dataset": it["dataset"],
                       "seconds": round(time.time() - t0, 2), **out}
                with lock:
                    with open(path, "a") as f:
                        f.write(json.dumps(row, default=str) + "\n")
                    rows.append(row)
            todo = [it for it in items if it["item_id"] not in done]
            with ThreadPoolExecutor(max_workers=max(1, a.workers)) as ex:
                list(ex.map(work, todo))
            summary[f"{m}__{jn}"] = score(rows)
            by_rel = defaultdict(list)
            for r in rows:
                by_rel[r["relation"]].append(r)
            summary[f"{m}__{jn}"]["by_relation"] = {k: score(v) for k, v in sorted(by_rel.items())}
            print(f"{m:20} {jn:12}", json.dumps({k: v for k, v in summary[f'{m}__{jn}'].items() if k != 'by_relation'}))
    if ctx:
        ctx.__exit__(None, None, None)
    if a.live:
        once = MECH / "eval_v3_once.json"
        prev = json.load(open(once)) if once.exists() else {}
        for k, v in summary.items():
            prev.setdefault(k, v)          # the first complete live score of a method is kept, never replaced
        prev["note"] = ("held-out v3: annotator-written sentences, checker executable/1.4 frozen before scoring; "
                        "writer-intent labels, blind second-rater validation pending")
        once.write_text(json.dumps(prev, indent=1))
    return 0


if __name__ == "__main__":
    import os
    rc = main()
    sys.stdout.flush()
    os._exit(rc)
