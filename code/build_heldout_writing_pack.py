"""
Annotator-written held-out mechanism set (Future_Plan §E2; Research_Plan §9.5 protocol item).

The v1/v2 mechanism packs were worded by the method developers. This pack asks an INDEPENDENT
annotator to write the four versions of each base claim in their own words, from the evidence only:

  original     a correct one-sentence statement of the assigned relation
  paraphrase   the same correct content, worded differently
  corruption   the same relation with exactly ONE genuine error of that relation's type
  provocation  a CORRECT sentence that a careless checker might wrongly object to

Borrowers: development-role borrowers never used in M1, practice, mechanism v1 or v2. The
annotator never sees the checker, its rules beyond plain definitions, or any earlier template.

    .venv/bin/python build_heldout_writing_pack.py --per-dataset 15
Output: Submission/packs/Heldout_writing_pack.zip  (to the annotator)
        Submission/mechanism/annotation_key_mechanism_v3.jsonl  (LOCAL ONLY: borrower ids, evidence)
"""
from __future__ import annotations

import argparse
import json
import random
import shutil
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_mechanism_pack import BUILDERS, OUT as MECH, _attrs, used_borrowers  # noqa: E402
from build_pack_m1 import SUB, assert_blind  # noqa: E402
from evidence_builder import render_evidence_v2  # noqa: E402

PACKS = SUB / "packs"
SEED = 20261009
RELS_TW = ["VALUE", "DIRECTION", "RANK", "SCALE", "TEMPORAL", "CODE", "CAUSAL", "PD"]
RELS_HC = ["VALUE", "DIRECTION", "RANK", "SCALE", "CAUSAL", "PD"]

REL_DEF = {
    "VALUE": "the recorded VALUE of the target field (number and unit as in the evidence; sensible rounding is fine)",
    "DIRECTION": "whether the target driver RAISES or LOWERS the model's risk score",
    "RANK": "which driver is the STRONGEST (largest contribution in absolute size)",
    "SCALE": "the size of the target driver's contribution, which is on the raw log-odds scale "
             "(it is NOT a percentage and NOT a change in probability)",
    "TEMPORAL": "how the target series CHANGED between the two dated observations shown",
    "CODE": "what the target CODE means - here the source does NOT document the code's meaning",
    "CAUSAL": "the relation between the target driver and the model's score, which is an ASSOCIATION in the model, "
              "not a real-world cause and not the effect of any action",
    "PD": "the model's estimated PROBABILITY OF DEFAULT and the outcome HORIZON it refers to",
}
ERR_HINT = {
    "VALUE": "a wrong value for that field (or another field's value put on this field's name)",
    "DIRECTION": "the opposite direction",
    "RANK": "a driver that is NOT the strongest called the strongest",
    "SCALE": "the contribution presented as a percentage or as a change in the probability of default",
    "TEMPORAL": "the wrong direction of change, or a change asserted where the evidence shows only one observation",
    "CODE": "a definite meaning given to the undocumented code",
    "CAUSAL": "a real-world causal claim or the claim that an action would change the default outcome",
    "PD": "a wrong probability, or a wrong horizon (e.g. 12 months when the evidence says otherwise)",
}
PROV_HINT = {
    "VALUE": "a correctly rounded or approximate value, or the value in a different but equivalent unit",
    "DIRECTION": "a correct direction stated with a negation or a contrast (e.g. 'does not ... ; it ...')",
    "RANK": "a correct ranking with a restricted scope (e.g. strongest among the risk-reducing drivers), or a "
            "correct mention of a non-strongest driver that also names the strongest",
    "SCALE": "a correct log-odds statement that MENTIONS percentages or probability only to deny that reading",
    "TEMPORAL": "a correct, supported change described with words like 'over time' or 'has risen'",
    "CODE": "a sentence that mentions a possible meaning only to say the source does not document it",
    "CAUSAL": "a disclaimer that uses causal words to DENY causation",
    "PD": "the correct probability with a correctly DENIED wrong horizon (e.g. 'not a 12-month estimate')",
}

README = """# Held-out sentence-writing task

Thank you for helping. You will write short sentences about credit-model evidence. Your sentences
become a test set for software that checks explanations. **The software and its rules are hidden
from you on purpose: write naturally, in your own words.**

## What you receive
- `Heldout_writing_workbook.xlsx`
  - `Evidence`: for each borrower (B01, B02, ...), the evidence: recorded facts, the model's drivers
    (each RAISES or LOWERS the model's risk score, with a contribution on the raw log-odds scale) and
    the estimated probability of default with its outcome and horizon.
  - `Write`: one row per sentence to write. Each row names the borrower, the RELATION to write about,
    the TARGET evidence item(s) and the VERSION.
- This README.

## The four versions (for every base claim)
1. **original** - one correct sentence about the relation.
2. **paraphrase** - the SAME correct content in clearly different words (different sentence
   structure, and if you like a different name for the field: its code, its plain-language name or
   a description).
3. **corruption** - the same relation with **exactly one** genuine error of the kind described in
   the row; everything else in the sentence must be correct. Describe the error in `INTENDED_ERROR`.
4. **provocation** - a **correct** sentence that might *look* wrong to a careless reader (a hint is
   given in the row). Describe the false objection it invites in `INVITED_OBJECTION`.

## Rules
- One sentence per row, at most about 40 words. Use only the borrower's evidence.
- Vary your wording: do not use the same sentence frame more than three times in the whole sheet.
  Use a mix of field codes (e.g. `PAY_AMT3`) and plain names ("the July payment").
- Do not copy whole phrases from the evidence text; names and numbers are fine.
- Corruptions contain ONE error; originals, paraphrases and provocations contain NONE.
- Contributions are log-odds numbers, never percentages. The probability of default is the only
  probability in the evidence.
- Undocumented codes have no known meaning.
- Do not use ChatGPT or any other AI tool, and do not search for template sentences.
- Fill `annotator_code`, `date`, `hours` and `used_AI_tool (must be N)` in the `Write` sheet header rows.

## Example (a made-up field, not in your evidence)
Evidence: `[a9] Driver: ACCOUNT_AGE LOWERS the model's risk score (contribution -0.21 on the raw log-odds margin)`.
Relation DIRECTION, target [a9]:
- original: "ACCOUNT_AGE pulls the model's risk score down."
- paraphrase: "Compared with the model's reference, the age of the account works in the applicant's favour."
- corruption: "The account's age adds to the model's risk." (INTENDED_ERROR: direction reversed)
- provocation: "Account age is not what raises the score here; it lowers it." (INVITED_OBJECTION: a
  checker may read 'raises the score' and flag a reversal)

Estimated time: 5-6 hours for 240 sentences. Questions: ask the study author, not other annotators.
"""


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from build_pack_m1 import extra_borrowers

    ap = argparse.ArgumentParser()
    ap.add_argument("--per-dataset", type=int, default=15)
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()
    keyp = MECH / "annotation_key_mechanism_v3.jsonl"
    if keyp.exists() and not a.force:
        print(f"{keyp} exists: the v3 borrower set is fixed once built")
        return 1
    rng = random.Random(SEED)
    exclude = used_borrowers()
    for v in ("v1", "v2"):
        p = MECH / f"mechanism_pack_{v}.jsonl"
        if p.exists():
            exclude |= {json.loads(l)["borrower_id"] for l in open(p)}

    rows, key, evid = [], [], []
    k = 0
    for ds, rels in (("taiwan", RELS_TW), ("home_credit", RELS_HC)):
        made = 0
        for eo, _ in extra_borrowers(ds, exclude, a.per_dataset * 3, rng):
            if made >= a.per_dataset:
                break
            got = []
            for _ in range(len(rels)):
                rel = rels[k % len(rels)]
                k += 1
                at = _attrs(eo)
                cands = [at[0]] if rel in ("RANK", "PD", "CODE") else list(at)
                rng.shuffle(cands) if rel not in ("RANK", "PD", "CODE") else None
                out = None
                for c in cands:
                    out = BUILDERS[rel](eo, c, "dev", random.Random(0))
                    if out:
                        break
                if out and rel not in [g[0] for g in got]:
                    got.append((rel, out["_ids"]))
                if len(got) == 2:
                    break
            if len(got) < 2:
                continue
            made += 1
            bcode = f"B{len(evid) + 1:02d}"
            evid.append((bcode, ds, render_evidence_v2(eo)))
            for j, (rel, ids) in enumerate(got):
                base = f"HW-{bcode}-{'AB'[j]}"
                for ver in ("original", "paraphrase", "corruption", "provocation"):
                    what = {"original": f"A correct sentence about {REL_DEF[rel]}.",
                            "paraphrase": "The same correct content as your 'original', in clearly different words.",
                            "corruption": f"Exactly one error: {ERR_HINT[rel]}.",
                            "provocation": f"A CORRECT sentence that may look wrong, e.g. {PROV_HINT[rel]}."}[ver]
                    rows.append([f"{base}-{ver[:4]}", bcode, ds.replace("_", " "), rel, ", ".join(ids), ver, what,
                                 "", "", ""])
            key.append({"borrower_code": bcode, "borrower_id": eo.borrower_id, "dataset": ds,
                        "relations": [{"relation": r, "target_ids": i} for r, i in got],
                        "evidence": eo.to_dict(include_groups=True)})

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "README"
    for line in README.splitlines():
        ws0.append([line])
    ws0.column_dimensions["A"].width = 120
    we = wb.create_sheet("Evidence")
    we.append(["borrower_code", "dataset", "evidence"])
    for b, ds, ev in evid:
        we.append([b, ds.replace("_", " "), ev])
    we.column_dimensions["C"].width = 140
    ww = wb.create_sheet("Write")
    ww.append(["annotator_code", "", "date", "", "hours", "", "used_AI_tool (must be N)", ""])
    ww.append([])
    hdr = ["item_code", "borrower_code", "dataset", "relation", "target_evidence", "version", "what_to_write",
           "TEXT", "INTENDED_ERROR (corruption only)", "INVITED_OBJECTION (provocation only)"]
    ww.append(hdr)
    for r in rows:
        ww.append(r)
    bold, head = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    for cell in ww[3]:
        cell.font, cell.fill = bold, head
    for ws in (we, ww):
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = wrap
    for col, w in zip("ABCDEFGHIJ", (16, 8, 12, 11, 14, 12, 50, 60, 30, 30)):
        ww.column_dimensions[col].width = w
    for row in ww.iter_rows(min_row=4, min_col=8, max_col=10):
        for cell in row:
            cell.fill = fill
    ww.freeze_panes = "H4"
    bids = [k_["borrower_id"] for k_ in key]
    for ws in (we, ww):
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert_blind(cell.value, bids)

    stage = PACKS / "Heldout_writing_pack"
    shutil.rmtree(stage, ignore_errors=True)
    stage.mkdir(parents=True)
    wb.save(stage / "Heldout_writing_workbook.xlsx")
    (stage / "README_writing_task.md").write_text(README)
    z = PACKS / "Heldout_writing_pack.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(stage.iterdir()):
            zf.write(f, f"Heldout_writing_pack/{f.name}")
    with open(keyp, "w") as f:
        for kk in key:
            f.write(json.dumps(kk, default=str) + "\n")
    from collections import Counter
    print(z, f"{len(evid)} borrowers, {len(rows) // 4} base claims, {len(rows)} sentences;",
          dict(Counter(r[3] for r in rows if r[5] == "original")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
